"""
produk_excel.py
===============
Fitur Import / Export Excel untuk halaman Produk (YN Hijab).

Endpoint (semua di bawah /produk/excel):
    GET  /produk/excel/template  -> unduh template Excel (.xlsx)
    POST /produk/excel/import    -> unggah .xlsx, hasil berupa JSON
    GET  /produk/excel/export    -> unduh data produk dari database (.xlsx)

Aturan data mengikuti proses "Tambah Produk" di app.py:
    - produk       : id_kategori, kode_produk, nama_produk, harga_jual,
                     gambar (NULL), status 'Aktif'
    - varian_produk: satu baris per warna
    - stok_barang  : satu baris per varian (id_rak, stok awal, stok_minimum 5)
Import HANYA menambah produk baru. Produk/stok yang sudah ada tidak diubah.
"""

import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from models.database import get_db_connection


produk_excel_bp = Blueprint(
    "produk_excel",
    __name__,
    url_prefix="/produk/excel"
)


# =========================================================
# KONSTANTA
# =========================================================

SHEET_DATA = "Data Produk"
SHEET_CONTOH = "Contoh Pengisian"
SHEET_PETUNJUK = "Petunjuk"
SHEET_REFERENSI = "Referensi"

# Urutan kolom template / kolom yang dibaca saat import
KOLOM_IMPOR = [
    "Kode Produk",
    "Nama Produk",
    "Kategori",
    "Warna",
    "Harga Jual",
    "Rak",
    "Stok Awal",
]
KOLOM_WAJIB = ["Nama Produk", "Kategori", "Warna", "Harga Jual", "Rak"]

KODE_PREFIX = "PR"          # sama dengan tambah_produk() di app.py
STOK_MINIMUM_DEFAULT = 5    # sama dengan tambah_produk() di app.py
STATUS_AKTIF = "Aktif"

MAKS_UKURAN_FILE = 5 * 1024 * 1024    # 5 MB
MAKS_BARIS_DATA = 2000
MAKS_PANJANG_KODE = 20                # produk.kode_produk varchar(20)
MAKS_PANJANG_NAMA = 100               # produk.nama_produk varchar(100)
HARGA_MAKS = Decimal("9999999999.99")  # produk.harga_jual decimal(12,2)
STOK_MAKS = 2147483647                # int(11)
MAKS_PESAN_GAGAL = 300                # batas baris gagal yang dikirim ke browser

MIMETYPE_XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


# =========================================================
# STYLE EXCEL
# =========================================================

WARNA_HEADER_WAJIB = "7C2D52"
WARNA_HEADER_OPSIONAL = "6B7280"
WARNA_HEADER_EXPORT = "7C2D52"
WARNA_BAND = "F7F1F4"

FONT_HEADER = Font(bold=True, color="FFFFFF")
GARIS_TIPIS = Side(style="thin", color="D1D5DB")
GARIS_TEBAL = Side(style="medium", color="7C2D52")


def _isi_header(ws, kolom, wajib=None, baris=1):
    """Tulis baris header dengan style."""
    for no, nama in enumerate(kolom, start=1):
        sel = ws.cell(row=baris, column=no, value=nama)
        if wajib is None or nama in wajib:
            warna = WARNA_HEADER_WAJIB
        else:
            warna = WARNA_HEADER_OPSIONAL
        sel.fill = PatternFill("solid", fgColor=warna)
        sel.font = FONT_HEADER
        sel.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[baris].height = 24


def _set_teks(sel, nilai):
    """
    Isi sel dengan teks. Teks yang diawali '=' dipaksa menjadi string
    supaya tidak dianggap rumus oleh Excel.
    """
    sel.value = nilai
    if isinstance(nilai, str) and nilai.startswith("="):
        sel.data_type = "s"


def _atur_lebar(ws, lebar):
    for no, nilai in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(no)].width = nilai


# =========================================================
# AKSES DATABASE (REFERENSI)
# =========================================================

def _norm(teks):
    """Normalisasi untuk pencocokan: huruf kecil, spasi dirapikan."""
    if teks is None:
        return ""
    return re.sub(r"\s+", " ", str(teks).replace("*", "")).strip().lower()


def _muat_referensi(cursor):
    """
    Ambil kategori, warna, rak, dan kode produk yang sudah ada.
    Semua dibaca dari tabel asli: kategori, warna, rak, produk.
    """
    cursor.execute("""
        SELECT id_kategori, nama_kategori
        FROM kategori
        ORDER BY nama_kategori
    """)
    kategori = cursor.fetchall()

    cursor.execute("""
        SELECT id_warna, nama_warna
        FROM warna
        ORDER BY nama_warna
    """)
    warna = cursor.fetchall()

    cursor.execute("""
        SELECT id_rak, kode_rak, nama_rak
        FROM rak
        ORDER BY kode_rak
    """)
    rak = cursor.fetchall()

    cursor.execute("SELECT kode_produk FROM produk")
    kode_ada = {_norm(r["kode_produk"]) for r in cursor.fetchall()}

    cursor.execute("""
        SELECT COALESCE(
                   MAX(CAST(SUBSTRING(kode_produk, 3) AS UNSIGNED)),
                   0
               ) AS nomor_terakhir
        FROM produk
        WHERE kode_produk REGEXP '^PR[0-9]+$'
    """)
    nomor_terakhir = int(cursor.fetchone()["nomor_terakhir"] or 0)

    peta_rak = {}
    for r in rak:
        peta_rak[_norm(r["kode_rak"])] = r["id_rak"]
        peta_rak.setdefault(_norm(r["nama_rak"]), r["id_rak"])

    return {
        "daftar_kategori": kategori,
        "daftar_warna": warna,
        "daftar_rak": rak,
        "kategori": {_norm(k["nama_kategori"]): k["id_kategori"]
                     for k in kategori},
        "warna": {_norm(w["nama_warna"]): (w["id_warna"], w["nama_warna"])
                  for w in warna},
        "rak": peta_rak,
        "kode_ada": kode_ada,
        "nomor_terakhir": nomor_terakhir,
    }


# =========================================================
# PARSING NILAI
# =========================================================

def _teks(nilai):
    """Nilai sel -> string rapi. Angka bulat (5.0) menjadi '5'."""
    if nilai is None:
        return ""
    if isinstance(nilai, float) and nilai.is_integer():
        return str(int(nilai))
    return str(nilai).strip()


def _parse_angka(nilai):
    """
    Ubah nilai sel menjadi Decimal.
    Mendukung angka Excel dan teks seperti '170000', '170.000',
    'Rp 170.000', '170.000,50'. Melempar ValueError jika tidak valid.
    """
    if isinstance(nilai, bool):
        raise ValueError("bukan angka")

    if isinstance(nilai, (int, float, Decimal)):
        angka = Decimal(str(nilai))
    else:
        teks = str(nilai).strip().lower().replace("rp", "")
        teks = teks.replace(" ", "")
        if not teks:
            raise ValueError("kosong")

        if "," in teks and "." in teks:
            teks = teks.replace(".", "").replace(",", ".")
        elif "." in teks and re.fullmatch(r"-?\d{1,3}(\.\d{3})+", teks):
            teks = teks.replace(".", "")
        elif "," in teks:
            teks = teks.replace(",", ".")

        try:
            angka = Decimal(teks)
        except InvalidOperation:
            raise ValueError("bukan angka")

    if not angka.is_finite():
        raise ValueError("bukan angka")

    return angka


def _parse_harga(nilai):
    try:
        angka = _parse_angka(nilai)
    except ValueError:
        raise ValueError("Harga Jual harus berupa angka.")

    if angka < 0:
        raise ValueError("Harga Jual tidak boleh negatif.")

    if angka != angka.quantize(Decimal("0.01")):
        raise ValueError("Harga Jual maksimal 2 angka di belakang koma.")

    angka = angka.quantize(Decimal("0.01"))

    if angka > HARGA_MAKS:
        raise ValueError("Harga Jual terlalu besar (maksimal 9.999.999.999).")

    return angka


def _parse_stok(nilai):
    """Stok awal: kosong = 0. Harus bilangan bulat >= 0."""
    if nilai is None or _teks(nilai) == "":
        return 0

    try:
        angka = _parse_angka(nilai)
    except ValueError:
        raise ValueError("Stok Awal harus berupa angka bulat.")

    if angka != angka.to_integral_value():
        raise ValueError("Stok Awal harus berupa bilangan bulat.")

    angka = int(angka)

    if angka < 0:
        raise ValueError("Stok Awal tidak boleh negatif.")

    if angka > STOK_MAKS:
        raise ValueError("Stok Awal terlalu besar.")

    return angka


# =========================================================
# MEMBACA FILE EXCEL
# =========================================================

class ErrorFile(Exception):
    """Kesalahan pada file secara keseluruhan (bukan per baris)."""


def _baca_baris_excel(file_storage):
    """
    Baca file .xlsx yang diunggah.
    Mengembalikan list dict: {"baris", "kode", "nama", "kategori",
    "warna", "harga", "rak", "stok"}. Baris kosong dilewati.
    Melempar ErrorFile bila file tidak layak dibaca.
    """
    nama_file = (file_storage.filename or "").strip()

    if not nama_file.lower().endswith(".xlsx"):
        raise ErrorFile("File harus berformat .xlsx.")

    data = file_storage.read(MAKS_UKURAN_FILE + 1)

    if len(data) > MAKS_UKURAN_FILE:
        raise ErrorFile("Ukuran file terlalu besar (maksimal 5 MB).")

    if not data.startswith(b"PK"):
        raise ErrorFile("File bukan Excel .xlsx yang valid.")

    try:
        wb = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=True
        )
    except Exception:
        raise ErrorFile("File Excel tidak dapat dibaca atau rusak.")

    try:
        if SHEET_DATA in wb.sheetnames:
            ws = wb[SHEET_DATA]
        else:
            ws = wb.worksheets[0]

        baris_iter = ws.iter_rows(values_only=True)

        header = next(baris_iter, None)

        if not header:
            raise ErrorFile("Sheet kosong. Gunakan template dari tombol "
                            "'Template Excel'.")

        posisi = {}
        for idx, nama in enumerate(header):
            kunci = _norm(nama)
            if kunci and kunci not in posisi:
                posisi[kunci] = idx

        kurang = [k for k in KOLOM_WAJIB if _norm(k) not in posisi]
        if kurang:
            raise ErrorFile(
                "Kolom wajib tidak ditemukan: " + ", ".join(kurang) +
                ". Gunakan template dari tombol 'Template Excel'."
            )

        def ambil(baris, nama_kolom):
            idx = posisi.get(_norm(nama_kolom))
            if idx is None or idx >= len(baris):
                return None
            return baris[idx]

        hasil = []
        kosong_berurutan = 0

        for nomor, baris in enumerate(baris_iter, start=2):
            item = {
                "baris": nomor,
                "kode": ambil(baris, "Kode Produk"),
                "nama": ambil(baris, "Nama Produk"),
                "kategori": ambil(baris, "Kategori"),
                "warna": ambil(baris, "Warna"),
                "harga": ambil(baris, "Harga Jual"),
                "rak": ambil(baris, "Rak"),
                "stok": ambil(baris, "Stok Awal"),
            }

            kosong = all(
                _teks(v) == ""
                for k, v in item.items() if k != "baris"
            )

            if kosong:
                kosong_berurutan += 1
                if kosong_berurutan >= 200:
                    break
                continue

            kosong_berurutan = 0
            hasil.append(item)

            if len(hasil) > MAKS_BARIS_DATA:
                raise ErrorFile(
                    f"Jumlah baris data melebihi {MAKS_BARIS_DATA}. "
                    "Pecah menjadi beberapa file."
                )

        return hasil

    finally:
        wb.close()


# =========================================================
# VALIDASI
# =========================================================

def _validasi(baris_excel, ref):
    """
    Validasi semua baris dan kelompokkan menjadi produk.

    Pengelompokan (satu produk dapat punya banyak warna):
      - Jika Kode Produk diisi  -> baris dengan kode sama = satu produk.
      - Jika Kode Produk kosong -> baris dengan Nama + Kategori sama
                                   = satu produk.

    Jika ada satu saja baris bermasalah pada sebuah produk, SELURUH
    produk itu tidak disimpan (tidak ada produk setengah tersimpan).

    Mengembalikan (produk_valid, gagal):
      produk_valid : list produk siap simpan
      gagal        : list {"baris": n, "alasan": "..."} terurut
    """
    kelompok = {}        # kunci -> dict kelompok (urutan masuk dipertahankan)
    pesan_baris = {}     # nomor baris -> list pesan

    def catat(nomor, pesan):
        pesan_baris.setdefault(nomor, []).append(pesan)

    for item in baris_excel:
        nomor = item["baris"]

        kode = _teks(item["kode"])
        nama = _teks(item["nama"])
        nama_kategori = _teks(item["kategori"])
        nama_warna = _teks(item["warna"])
        nama_rak = _teks(item["rak"])

        # ---------- validasi per kolom ----------

        if kode:
            if len(kode) > MAKS_PANJANG_KODE:
                catat(nomor, f"Kode Produk maksimal {MAKS_PANJANG_KODE} karakter.")
            elif not re.fullmatch(r"[A-Za-z0-9_\-]+", kode):
                catat(nomor, "Kode Produk hanya boleh huruf, angka, '-' dan '_'.")
            elif _norm(kode) in ref["kode_ada"]:
                catat(nomor, f"Kode Produk '{kode}' sudah dipakai produk lain.")

        if not nama:
            catat(nomor, "Nama Produk wajib diisi.")
        elif len(nama) > MAKS_PANJANG_NAMA:
            catat(nomor, f"Nama Produk maksimal {MAKS_PANJANG_NAMA} karakter.")

        id_kategori = None
        if not nama_kategori:
            catat(nomor, "Kategori wajib diisi.")
        else:
            id_kategori = ref["kategori"].get(_norm(nama_kategori))
            if id_kategori is None:
                catat(nomor, f"Kategori '{nama_kategori}' tidak ditemukan.")

        id_warna = None
        if not nama_warna:
            catat(nomor, "Warna wajib diisi.")
        else:
            data_warna = ref["warna"].get(_norm(nama_warna))
            if data_warna is None:
                catat(nomor, f"Warna '{nama_warna}' tidak ditemukan.")
            else:
                id_warna = data_warna[0]

        harga = None
        if _teks(item["harga"]) == "":
            catat(nomor, "Harga Jual wajib diisi.")
        else:
            try:
                harga = _parse_harga(item["harga"])
            except ValueError as e:
                catat(nomor, str(e))

        id_rak = None
        if not nama_rak:
            catat(nomor, "Rak wajib diisi.")
        else:
            id_rak = ref["rak"].get(_norm(nama_rak))
            if id_rak is None:
                catat(nomor, f"Rak '{nama_rak}' tidak ditemukan.")

        stok = 0
        try:
            stok = _parse_stok(item["stok"])
        except ValueError as e:
            catat(nomor, str(e))

        # ---------- pengelompokan ----------

        if kode:
            kunci = ("kode", _norm(kode))
        elif nama:
            kunci = ("nama", _norm(nama), _norm(nama_kategori))
        else:
            kunci = ("baris", nomor)

        grup = kelompok.get(kunci)

        if grup is None:
            grup = {
                "kode": kode,
                "nama": nama,
                "kategori_teks": nama_kategori,
                "id_kategori": id_kategori,
                "harga": harga,
                "baris_pertama": nomor,
                "baris": [],
                "varian": [],
                "warna_dipakai": {},
            }
            kelompok[kunci] = grup
        else:
            # Konsistensi dengan baris pertama produk yang sama
            if kode and nama and _norm(nama) != _norm(grup["nama"]):
                catat(nomor, f"Nama Produk berbeda dari baris "
                             f"{grup['baris_pertama']} untuk kode yang sama.")
            if kode and nama_kategori and \
                    _norm(nama_kategori) != _norm(grup["kategori_teks"]):
                catat(nomor, f"Kategori berbeda dari baris "
                             f"{grup['baris_pertama']} untuk kode yang sama.")
            if harga is not None and grup["harga"] is not None \
                    and harga != grup["harga"]:
                catat(nomor, f"Harga Jual berbeda dari baris "
                             f"{grup['baris_pertama']} untuk produk yang sama.")

            # Lengkapi data grup bila baris pertama kosong/invalid
            if grup["harga"] is None and harga is not None:
                grup["harga"] = harga
            if grup["id_kategori"] is None and id_kategori is not None:
                grup["id_kategori"] = id_kategori
            if not grup["nama"] and nama:
                grup["nama"] = nama

        if id_warna is not None:
            if id_warna in grup["warna_dipakai"]:
                catat(nomor, f"Warna '{nama_warna}' sudah ada pada produk "
                             f"yang sama (baris "
                             f"{grup['warna_dipakai'][id_warna]}).")
            else:
                grup["warna_dipakai"][id_warna] = nomor

        grup["baris"].append(nomor)
        grup["varian"].append({
            "baris": nomor,
            "id_warna": id_warna,
            "id_rak": id_rak,
            "stok": stok,
        })

    # ---------- keputusan per produk ----------

    produk_valid = []
    gagal = {}

    for grup in kelompok.values():
        baris_bermasalah = [n for n in grup["baris"] if n in pesan_baris]

        if baris_bermasalah:
            for n in grup["baris"]:
                if n not in pesan_baris:
                    gagal[n] = (
                        "Dilewati: produk ini memiliki baris bermasalah "
                        "(baris " +
                        ", ".join(str(b) for b in baris_bermasalah) +
                        "), sehingga seluruh produk tidak disimpan."
                    )
            continue

        produk_valid.append(grup)

    for nomor, pesan in pesan_baris.items():
        gagal[nomor] = " ".join(pesan)

    daftar_gagal = [
        {"baris": n, "alasan": gagal[n]} for n in sorted(gagal)
    ]

    return produk_valid, daftar_gagal


# =========================================================
# MENYIMPAN KE DATABASE
# =========================================================

def _simpan_produk(cursor, produk_valid, ref):
    """
    Simpan produk baru. Pola sama dengan tambah_produk() di app.py:
    produk -> varian_produk -> stok_barang.
    Tidak menyentuh produk/stok yang sudah ada.
    Commit/rollback dilakukan oleh pemanggil.

    Mengembalikan (jumlah_produk, jumlah_varian).
    """
    kode_terpakai = set(ref["kode_ada"])

    for p in produk_valid:
        if p["kode"]:
            kode_terpakai.add(_norm(p["kode"]))

    nomor = ref["nomor_terakhir"]
    jumlah_varian = 0

    for p in produk_valid:

        if p["kode"]:
            kode_produk = p["kode"]
        else:
            while True:
                nomor += 1
                kode_produk = f"{KODE_PREFIX}{nomor:03d}"
                if _norm(kode_produk) not in kode_terpakai:
                    break
            kode_terpakai.add(_norm(kode_produk))

        cursor.execute("""
            INSERT INTO produk
            (
                id_kategori,
                kode_produk,
                nama_produk,
                harga_jual,
                gambar,
                status
            )
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (
            p["id_kategori"],
            kode_produk,
            p["nama"],
            p["harga"],
            None,
            STATUS_AKTIF
        ))

        id_produk = cursor.lastrowid

        for v in p["varian"]:

            cursor.execute("""
                INSERT INTO varian_produk
                (id_produk, id_warna)
                VALUES (%s, %s)
            """, (
                id_produk,
                v["id_warna"]
            ))

            id_varian = cursor.lastrowid

            cursor.execute("""
                INSERT INTO stok_barang
                (
                    id_varian,
                    id_rak,
                    stok,
                    stok_minimum
                )
                VALUES (%s, %s, %s, %s)
            """, (
                id_varian,
                v["id_rak"],
                v["stok"],
                STOK_MINIMUM_DEFAULT
            ))

            jumlah_varian += 1

    return len(produk_valid), jumlah_varian


# =========================================================
# ROUTE: DOWNLOAD TEMPLATE
# =========================================================

def _isi_sheet_referensi(ws, ref):
    ws["A1"], ws["C1"], ws["E1"], ws["F1"] = (
        "Kategori", "Warna", "Kode Rak", "Nama Rak"
    )
    for sel in (ws["A1"], ws["C1"], ws["E1"], ws["F1"]):
        sel.fill = PatternFill("solid", fgColor=WARNA_HEADER_OPSIONAL)
        sel.font = FONT_HEADER

    for i, k in enumerate(ref["daftar_kategori"], start=2):
        _set_teks(ws.cell(row=i, column=1), k["nama_kategori"])
    for i, w in enumerate(ref["daftar_warna"], start=2):
        _set_teks(ws.cell(row=i, column=3), w["nama_warna"])
    for i, r in enumerate(ref["daftar_rak"], start=2):
        _set_teks(ws.cell(row=i, column=5), r["kode_rak"])
        _set_teks(ws.cell(row=i, column=6), r["nama_rak"])

    ws["H1"] = ("Sheet ini berisi daftar pilihan dari database. "
                "Jangan diubah atau dihapus.")
    ws["H1"].font = Font(italic=True, color="6B7280")

    _atur_lebar(ws, [24, 3, 22, 3, 12, 30, 3, 60])


def _tambah_dropdown(ws, ref):
    """Dropdown & validasi angka pada sheet Data Produk."""
    akhir = MAKS_BARIS_DATA + 1

    def dropdown(kolom, rentang):
        dv = DataValidation(
            type="list",
            formula1=rentang,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Pilihan tidak valid",
            error="Pilih nilai dari daftar (lihat sheet Referensi).",
        )
        dv.add(f"{kolom}2:{kolom}{akhir}")
        ws.add_data_validation(dv)

    if ref["daftar_kategori"]:
        n = len(ref["daftar_kategori"]) + 1
        dropdown("C", f"={SHEET_REFERENSI}!$A$2:$A${n}")
    if ref["daftar_warna"]:
        n = len(ref["daftar_warna"]) + 1
        dropdown("D", f"={SHEET_REFERENSI}!$C$2:$C${n}")
    if ref["daftar_rak"]:
        n = len(ref["daftar_rak"]) + 1
        dropdown("F", f"={SHEET_REFERENSI}!$E$2:$E${n}")

    dv_harga = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Harga tidak valid",
        error="Harga Jual harus berupa angka dan tidak boleh negatif.",
    )
    dv_harga.add(f"E2:E{akhir}")
    ws.add_data_validation(dv_harga)

    dv_stok = DataValidation(
        type="whole",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Stok tidak valid",
        error="Stok Awal harus bilangan bulat 0 atau lebih.",
    )
    dv_stok.add(f"G2:G{akhir}")
    ws.add_data_validation(dv_stok)


@produk_excel_bp.route("/template")
def download_template():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        ref = _muat_referensi(cursor)
    except Exception as e:
        print("ERROR TEMPLATE PRODUK:", e)
        return "Gagal membuat template: data referensi tidak dapat dibaca.", 500
    finally:
        cursor.close()
        conn.close()

    wb = Workbook()

    # ---------- Sheet 1: Data Produk (kosong, siap diisi) ----------
    ws = wb.active
    ws.title = SHEET_DATA
    _isi_header(ws, KOLOM_IMPOR, wajib=KOLOM_WAJIB)
    _atur_lebar(ws, [16, 32, 18, 20, 16, 12, 12])
    ws.freeze_panes = "A2"
    for baris in range(2, MAKS_BARIS_DATA + 2):
        ws.cell(row=baris, column=5).number_format = "#,##0"
    _tambah_dropdown(ws, ref)

    # ---------- Sheet 2: Contoh Pengisian ----------
    wc = wb.create_sheet(SHEET_CONTOH)
    wc.sheet_properties.tabColor = "9CA3AF"
    _isi_header(wc, KOLOM_IMPOR, wajib=KOLOM_WAJIB)
    _atur_lebar(wc, [16, 32, 18, 20, 16, 12, 12])

    kat = ref["daftar_kategori"][0]["nama_kategori"] \
        if ref["daftar_kategori"] else "Mukena"
    warna = [w["nama_warna"] for w in ref["daftar_warna"]] or \
        ["Black", "Navy"]
    while len(warna) < 2:
        warna.append(warna[0])
    rak = [r["kode_rak"] for r in ref["daftar_rak"]] or ["A01"]
    rak2 = rak[1] if len(rak) > 1 else rak[0]

    contoh = [
        ["", "Mukena Contoh Satu", kat, warna[0], 175000, rak[0], 10],
        ["", "Mukena Contoh Satu", kat, warna[1], 175000, rak[0], 8],
        ["", "Mukena Contoh Dua", kat, warna[0], 200000, rak2, 5],
        ["KD-001", "Mukena Contoh Tiga", kat, warna[1], 150000, rak[0], 0],
    ]
    for i, isi in enumerate(contoh, start=2):
        for j, nilai in enumerate(isi, start=1):
            sel = wc.cell(row=i, column=j)
            if nilai == "":
                continue
            _set_teks(sel, nilai)
        wc.cell(row=i, column=5).number_format = "#,##0"

    wc["A7"] = ("CONTOH SAJA - jangan diunggah. Isi data pada sheet "
                f"'{SHEET_DATA}'.")
    wc["A7"].font = Font(bold=True, color="B91C1C")

    # ---------- Sheet 3: Petunjuk ----------
    wp = wb.create_sheet(SHEET_PETUNJUK)
    petunjuk = [
        ("PETUNJUK PENGISIAN TEMPLATE PRODUK", True),
        ("", False),
        (f"1. Isi data pada sheet '{SHEET_DATA}'. Satu baris = satu warna "
         "(varian) dari sebuah produk.", False),
        ("2. Kolom berwarna merah tua wajib diisi: Nama Produk, Kategori, "
         "Warna, Harga Jual, Rak.", False),
        ("3. Kode Produk (opsional): jika kosong, sistem membuat otomatis "
         f"({KODE_PREFIX}001, {KODE_PREFIX}002, ...). Jika diisi, tidak "
         "boleh sama dengan kode produk yang sudah ada.", False),
        ("4. Produk dengan beberapa warna: tulis beberapa baris dengan Nama "
         "Produk + Kategori yang sama (atau Kode Produk yang sama) dan "
         "Harga Jual yang sama, warnanya berbeda.", False),
        ("5. Kategori, Warna, dan Rak harus sesuai daftar pada sheet "
         f"'{SHEET_REFERENSI}' (tersedia sebagai pilihan dropdown).", False),
        ("6. Harga Jual: angka tanpa 'Rp', tidak boleh negatif. "
         "Contoh: 175000.", False),
        ("7. Stok Awal: bilangan bulat 0 atau lebih; kosong dianggap 0. "
         "Stok awal berlaku untuk warna pada baris tersebut.", False),
        (f"8. Stok minimum otomatis {STOK_MINIMUM_DEFAULT} (sama seperti "
         "Tambah Produk).", False),
        ("9. Foto produk tidak diimpor; tambahkan foto lewat menu produk.", False),
        ("10. Import hanya MENAMBAH produk baru. Produk dan stok yang sudah "
         "ada tidak diubah atau ditimpa.", False),
        ("11. Jika ada baris bermasalah pada sebuah produk, seluruh produk "
         "itu tidak disimpan dan alasannya ditampilkan setelah import.", False),
        (f"12. Maksimal {MAKS_BARIS_DATA} baris dan ukuran file 5 MB.", False),
    ]
    for i, (teks, tebal) in enumerate(petunjuk, start=1):
        sel = wp.cell(row=i, column=1, value=teks)
        sel.alignment = Alignment(wrap_text=True, vertical="top")
        if tebal:
            sel.font = Font(bold=True, size=13)
    wp.column_dimensions["A"].width = 110

    # ---------- Sheet 4: Referensi ----------
    wr = wb.create_sheet(SHEET_REFERENSI)
    _isi_sheet_referensi(wr, ref)

    return _kirim_workbook(wb, "template_produk_yn_hijab.xlsx")


# =========================================================
# ROUTE: IMPORT
# =========================================================

@produk_excel_bp.route("/import", methods=["POST"])
def import_produk():
    file_excel = request.files.get("file_excel")

    if not file_excel or not file_excel.filename:
        return jsonify({
            "ok": False,
            "pesan": "Pilih file Excel (.xlsx) terlebih dahulu."
        }), 400

    try:
        baris_excel = _baca_baris_excel(file_excel)
    except ErrorFile as e:
        return jsonify({"ok": False, "pesan": str(e)}), 400

    if not baris_excel:
        return jsonify({
            "ok": False,
            "pesan": "Tidak ada baris data yang dapat diimpor."
        }), 400

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        ref = _muat_referensi(cursor)

        produk_valid, gagal = _validasi(baris_excel, ref)

        jumlah_produk = 0
        jumlah_varian = 0

        if produk_valid:
            jumlah_produk, jumlah_varian = _simpan_produk(
                cursor, produk_valid, ref
            )
            conn.commit()

    except Exception as e:
        conn.rollback()
        print("ERROR IMPORT PRODUK:", e)
        return jsonify({
            "ok": False,
            "pesan": ("Terjadi kesalahan pada database. "
                      "Tidak ada data yang tersimpan.")
        }), 500

    finally:
        cursor.close()
        conn.close()

    if jumlah_produk:
        pesan = (f"{jumlah_produk} produk ({jumlah_varian} varian warna) "
                 "berhasil diimpor.")
    else:
        pesan = "Tidak ada produk yang berhasil diimpor."

    if gagal:
        pesan += f" {len(gagal)} baris gagal atau dilewati."

    return jsonify({
        "ok": True,
        "pesan": pesan,
        "jumlah_produk": jumlah_produk,
        "jumlah_varian": jumlah_varian,
        "jumlah_gagal": len(gagal),
        "gagal": gagal[:MAKS_PESAN_GAGAL],
        "gagal_terpotong": len(gagal) > MAKS_PESAN_GAGAL,
    })


# =========================================================
# ROUTE: EXPORT
# =========================================================

KOLOM_EXPORT = [
    "Kode Produk",
    "Nama Produk",
    "Kategori",
    "Warna",
    "Harga Jual",
    "Stok",
    "Stok Minimum",
    "Status Stok",
    "Rak",
    "Status Produk",
]


@produk_excel_bp.route("/export")
def export_produk():
    # Default: hanya produk Aktif (sama dengan halaman Produk).
    # /produk/excel/export?semua=1 -> termasuk produk nonaktif.
    semua = request.args.get("semua") == "1"

    kondisi = "" if semua else "WHERE p.status = 'Aktif'"

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute(f"""
            SELECT p.id_produk,
                   p.kode_produk,
                   p.nama_produk,
                   p.harga_jual,
                   p.status,
                   k.nama_kategori,
                   w.nama_warna,
                   sb.stok,
                   sb.stok_minimum,
                   r.kode_rak
            FROM produk p
            LEFT JOIN kategori k
                ON p.id_kategori = k.id_kategori
            LEFT JOIN varian_produk vp
                ON p.id_produk = vp.id_produk
            LEFT JOIN warna w
                ON vp.id_warna = w.id_warna
            LEFT JOIN stok_barang sb
                ON vp.id_varian = sb.id_varian
            LEFT JOIN rak r
                ON sb.id_rak = r.id_rak
            {kondisi}
            ORDER BY p.kode_produk, p.id_produk, w.nama_warna
        """)
        data = cursor.fetchall()

    except Exception as e:
        print("ERROR EXPORT PRODUK:", e)
        return "Gagal mengekspor data produk.", 500

    finally:
        cursor.close()
        conn.close()

    wb = Workbook()

    # ---------- Sheet 1: Data Produk (satu baris per varian warna) ----------
    ws = wb.active
    ws.title = "Data Produk"
    _isi_header(ws, KOLOM_EXPORT, wajib=None)
    ws.freeze_panes = "C2"
    _atur_lebar(ws, [14, 32, 16, 20, 16, 10, 14, 14, 10, 14])

    ringkasan = {}      # id_produk -> dict
    urutan_produk = []
    id_sebelumnya = None
    band = False
    baris_excel = 1

    for r in data:
        baris_excel += 1

        if r["id_produk"] != id_sebelumnya:
            band = not band
            id_sebelumnya = r["id_produk"]
            awal_produk = True
        else:
            awal_produk = False

        ada_varian = r["nama_warna"] is not None
        stok = r["stok"]
        minimum = r["stok_minimum"]

        if not ada_varian or stok is None:
            status_stok = "-"
        elif minimum is not None and stok <= minimum:
            status_stok = "Menipis"
        else:
            status_stok = "Aman"

        status_produk = r["status"] if r["status"] else "Nonaktif"

        nilai = [
            r["kode_produk"],
            r["nama_produk"],
            r["nama_kategori"] or "-",
            r["nama_warna"] or "-",
            float(r["harga_jual"]) if r["harga_jual"] is not None else None,
            stok if stok is not None else "-",
            minimum if minimum is not None else "-",
            status_stok,
            r["kode_rak"] or "-",
            status_produk,
        ]

        for no, isi in enumerate(nilai, start=1):
            sel = ws.cell(row=baris_excel, column=no)
            _set_teks(sel, isi)

            if band:
                sel.fill = PatternFill("solid", fgColor=WARNA_BAND)

            sel.border = Border(
                top=GARIS_TEBAL if awal_produk else GARIS_TIPIS,
                bottom=GARIS_TIPIS,
            )

            if no == 5:
                sel.number_format = '"Rp" #,##0'
            if no in (6, 7, 8, 9, 10):
                sel.alignment = Alignment(horizontal="center")
            if no == 2 and awal_produk:
                sel.font = Font(bold=True)

        if status_stok == "Menipis":
            ws.cell(row=baris_excel, column=8).font = Font(
                bold=True, color="B91C1C"
            )

        # data untuk sheet ringkasan
        rg = ringkasan.get(r["id_produk"])
        if rg is None:
            rg = {
                "kode": r["kode_produk"],
                "nama": r["nama_produk"],
                "kategori": r["nama_kategori"] or "-",
                "harga": r["harga_jual"],
                "status": status_produk,
                "varian": 0,
                "stok": 0,
            }
            ringkasan[r["id_produk"]] = rg
            urutan_produk.append(r["id_produk"])
        if ada_varian:
            rg["varian"] += 1
            rg["stok"] += int(stok or 0)

    if data:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(KOLOM_EXPORT))}" \
                             f"{baris_excel}"

    # ---------- Sheet 2: Ringkasan Produk (satu baris per produk) ----------
    wr = wb.create_sheet("Ringkasan Produk")
    kolom_ringkas = [
        "Kode Produk", "Nama Produk", "Kategori", "Harga Jual",
        "Jumlah Varian", "Total Stok", "Status Produk",
    ]
    _isi_header(wr, kolom_ringkas, wajib=None)
    wr.freeze_panes = "A2"
    _atur_lebar(wr, [14, 32, 16, 16, 14, 12, 14])

    for i, id_produk in enumerate(urutan_produk, start=2):
        rg = ringkasan[id_produk]
        isi = [
            rg["kode"], rg["nama"], rg["kategori"],
            float(rg["harga"]) if rg["harga"] is not None else None,
            rg["varian"], rg["stok"], rg["status"],
        ]
        for no, nilai in enumerate(isi, start=1):
            sel = wr.cell(row=i, column=no)
            _set_teks(sel, nilai)
            if no == 4:
                sel.number_format = '"Rp" #,##0'
            if no in (5, 6, 7):
                sel.alignment = Alignment(horizontal="center")

    nama_file = "produk_yn_hijab_" + \
        datetime.now().strftime("%Y%m%d_%H%M") + ".xlsx"

    return _kirim_workbook(wb, nama_file)


# =========================================================
# UTILITAS
# =========================================================

def _kirim_workbook(wb, nama_file):
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=nama_file,
        mimetype=MIMETYPE_XLSX
    )