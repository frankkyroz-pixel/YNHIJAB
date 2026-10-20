"""
Fitur Import / Export Excel untuk halaman Barang Masuk - YN Hijab.

Letakkan file ini sejajar dengan app.py (sama seperti produk_excel.py),
lalu daftarkan blueprint-nya di app.py:

    from barang_masuk_excel import barang_masuk_excel_bp
    app.register_blueprint(barang_masuk_excel_bp)

Endpoint:
    GET  /barang-masuk/export-excel      -> unduh riwayat Barang Masuk (.xlsx)
    GET  /barang-masuk/template-excel    -> unduh template import (.xlsx)
    POST /barang-masuk/import/preview    -> validasi file + pratinjau (TIDAK menyimpan)
    POST /barang-masuk/import/konfirmasi -> simpan transaksi + tambah stok

Aturan penting yang mengikuti kode aplikasi yang sudah ada:
  * Satu baris Excel = satu transaksi barang_masuk (BM-xxxx) dengan satu
    baris detail_barang_masuk. Route edit/hapus yang sudah ada hanya
    memproses SATU detail per transaksi, jadi struktur ini wajib dijaga
    agar edit/hapus tetap benar.
  * Stok ditambah dengan UPDATE stok_barang SET stok = stok + jumlah
    (sama seperti tambah_barang_masuk). Rak tidak diubah: stok_barang
    berisi satu baris per varian dan rak sudah melekat di baris itu.
  * Semua penyimpanan berada dalam satu transaksi database.
"""

import hashlib
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from models.database import get_db_connection

barang_masuk_excel_bp = Blueprint("barang_masuk_excel", __name__)


# =========================================================
# KONFIGURASI
# =========================================================

MAX_FILE_BYTES = 2 * 1024 * 1024      # 2 MB
MAX_BARIS_DATA = 1000                  # maksimal baris data per file
MAX_BARIS_SCAN = 5000                  # batas pembacaan (termasuk baris kosong)
MAX_JUMLAH = 100000                    # batas wajar jumlah per baris
MAX_PREVIEW_BARIS = 300                # baris yang ditampilkan di pratinjau

NAMA_SHEET_IMPORT = "Import Barang Masuk"
NAMA_SHEET_PETUNJUK = "Petunjuk"
NAMA_SHEET_VARIAN = "Daftar Varian"
NAMA_SHEET_EXPORT = "Riwayat Barang Masuk"

# Penanda di kolom keterangan untuk mencegah file yang sama diproses dua kali.
# Tidak membuat tabel baru: penanda disimpan di barang_masuk.keterangan.
PENANDA_AWAL = "[IMPORT-EXCEL:"
POLA_PENANDA = re.compile(r"\s*\[IMPORT-EXCEL:[0-9a-f]+\]")

FONT = "Arial"
WARNA_HEADER = "D58A94"     # dusty pink, sama dengan tema aplikasi

# Kolom template import: (judul kolom, kunci internal)
KOLOM_IMPORT = [
    ("Tanggal", "tanggal"),
    ("Nama Penjahit", "penjahit"),
    ("Kode Produk", "kode"),
    ("Nama Produk (info)", "nama_produk"),
    ("Warna", "warna"),
    ("Jumlah", "jumlah"),
    ("Harga Beli per Pcs", "harga"),
    ("Keterangan", "keterangan"),
]

WAJIB = ["tanggal", "kode", "warna", "jumlah", "harga"]

# Nama judul kolom yang diterima (huruf kecil, spasi dirapikan)
ALIAS_KOLOM = {
    "tanggal": "tanggal",
    "tanggal masuk": "tanggal",
    "nama penjahit": "penjahit",
    "penjahit": "penjahit",
    "kode produk": "kode",
    "kode": "kode",
    "nama produk": "nama_produk",
    "nama produk (info)": "nama_produk",
    "warna": "warna",
    "nama warna": "warna",
    "jumlah": "jumlah",
    "jumlah masuk": "jumlah",
    "harga beli per pcs": "harga",
    "harga beli": "harga",
    "harga": "harga",
    "keterangan": "keterangan",
}

# Kolom khas file hasil EXPORT. File export bukan untuk di-import ulang
# karena akan menggandakan transaksi lama.
KOLOM_KHAS_EXPORT = {"no transaksi", "no faktur", "total"}


class FileError(Exception):
    """Kesalahan pada file secara keseluruhan (bukan pada satu baris)."""


# =========================================================
# UTILITAS UMUM
# =========================================================

def _norm(teks):
    """Normalisasi untuk pencocokan: huruf kecil + spasi dirapikan."""
    return re.sub(r"\s+", " ", str(teks or "")).strip().lower()


def _teks(nilai):
    """Ubah nilai sel menjadi teks bersih."""
    if nilai is None:
        return ""
    if isinstance(nilai, float) and nilai.is_integer():
        nilai = int(nilai)
    teks = str(nilai)
    teks = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", teks)
    return teks.strip()


def _teks_rapi(nilai):
    """Teks satu baris dengan spasi dirapikan (untuk kode, warna, nama)."""
    return re.sub(r"\s+", " ", _teks(nilai)).strip()


def _hapus_penanda(keterangan):
    return POLA_PENANDA.sub("", keterangan or "").strip()


def _set_teks(sel, nilai):
    """Isi sel dengan teks; teks berawalan '=' tidak boleh jadi rumus."""
    sel.value = nilai
    if isinstance(nilai, str) and nilai.startswith("="):
        sel.data_type = "s"


# =========================================================
# PARSING NILAI
# =========================================================

def _parse_tanggal(nilai):
    if isinstance(nilai, datetime):
        hasil = nilai.date()
    elif isinstance(nilai, date):
        hasil = nilai
    else:
        teks = _teks(nilai)
        if not teks:
            raise ValueError("Tanggal wajib diisi.")
        hasil = None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                hasil = datetime.strptime(teks, fmt).date()
                break
            except ValueError:
                continue
        if hasil is None:
            raise ValueError(
                f"Tanggal '{teks}' tidak valid. Gunakan format "
                f"YYYY-MM-DD, misalnya 2026-10-20."
            )

    if hasil.year < 2000 or hasil.year > date.today().year + 1:
        raise ValueError(f"Tanggal {hasil.isoformat()} tidak masuk akal.")
    return hasil


def _parse_jumlah(nilai):
    pesan_format = "Jumlah harus berupa bilangan bulat positif (contoh: 12)."

    if isinstance(nilai, bool):
        raise ValueError(pesan_format)
    if nilai is None or _teks(nilai) == "":
        raise ValueError("Jumlah wajib diisi.")

    if isinstance(nilai, int):
        angka = nilai
    elif isinstance(nilai, float):
        if not nilai.is_integer():
            raise ValueError(pesan_format)
        angka = int(nilai)
    else:
        teks = _teks(nilai)
        if not re.fullmatch(r"\d+", teks):
            raise ValueError(pesan_format)
        angka = int(teks)

    if angka <= 0:
        raise ValueError("Jumlah harus lebih dari 0.")
    if angka > MAX_JUMLAH:
        raise ValueError(f"Jumlah {angka} terlalu besar (maksimal {MAX_JUMLAH}).")
    return angka


def _parse_harga(nilai):
    pesan_format = (
        "Harga beli harus berupa angka tanpa pemisah ribuan "
        "(contoh: 150000)."
    )

    if isinstance(nilai, bool):
        raise ValueError(pesan_format)
    if nilai is None or _teks(nilai) == "":
        raise ValueError("Harga beli per pcs wajib diisi.")

    try:
        if isinstance(nilai, (int, float, Decimal)):
            angka = Decimal(str(nilai))
        else:
            teks = re.sub(r"(?i)^rp\.?\s*", "", _teks(nilai))
            if not re.fullmatch(r"\d+(\.\d{1,2})?", teks):
                raise ValueError(pesan_format)
            angka = Decimal(teks)
    except InvalidOperation:
        raise ValueError(pesan_format)

    if not angka.is_finite():
        raise ValueError(pesan_format)
    if angka < 0:
        raise ValueError("Harga beli tidak boleh negatif.")
    angka = angka.quantize(Decimal("0.01"))
    if angka >= Decimal("10000000000"):
        raise ValueError("Harga beli terlalu besar.")
    return angka


# =========================================================
# KATALOG VARIAN (dari database)
# =========================================================

def _muat_katalog(cursor):
    """Semua varian + info stok/rak. Kursor harus dictionary=True."""
    cursor.execute("""
        SELECT vp.id_varian,
               p.kode_produk,
               p.nama_produk,
               p.status AS status_produk,
               w.nama_warna,
               COUNT(sb.id_stok) AS jumlah_baris_stok,
               COALESCE(SUM(sb.stok), 0) AS stok,
               MIN(r.kode_rak) AS kode_rak
        FROM varian_produk vp
        JOIN produk p ON p.id_produk = vp.id_produk
        JOIN warna w ON w.id_warna = vp.id_warna
        LEFT JOIN stok_barang sb ON sb.id_varian = vp.id_varian
        LEFT JOIN rak r ON r.id_rak = sb.id_rak
        GROUP BY vp.id_varian, p.kode_produk, p.nama_produk,
                 p.status, w.nama_warna
        ORDER BY p.kode_produk, w.nama_warna
    """)
    daftar = cursor.fetchall()

    per_kunci = {}
    per_kode = {}
    for varian in daftar:
        kunci = (_norm(varian["kode_produk"]), _norm(varian["nama_warna"]))
        per_kunci.setdefault(kunci, []).append(varian)
        per_kode.setdefault(_norm(varian["kode_produk"]), []).append(varian)
    return daftar, per_kunci, per_kode


# =========================================================
# MEMBACA & MEMVALIDASI FILE
# =========================================================

def _ambil_file_upload():
    """Ambil bytes file dari request. Return (bytes, pesan_error)."""
    if request.content_length and request.content_length > MAX_FILE_BYTES + 512 * 1024:
        return None, "Ukuran file melebihi batas 2 MB."

    berkas = request.files.get("file")
    if not berkas or not berkas.filename:
        return None, "Pilih file Excel (.xlsx) terlebih dahulu."
    if not berkas.filename.lower().endswith(".xlsx"):
        return None, "Format file harus .xlsx (Excel). Format .xls atau .csv tidak didukung."

    isi = berkas.read(MAX_FILE_BYTES + 1)
    if len(isi) > MAX_FILE_BYTES:
        return None, "Ukuran file melebihi batas 2 MB."
    if not isi:
        return None, "File kosong."
    if not zipfile.is_zipfile(BytesIO(isi)):
        return None, "File bukan Excel .xlsx yang valid atau rusak."
    return isi, None


def _baca_sheet(isi):
    """Baca sheet import. Return (peta_kolom, daftar (nomor_baris, nilai))."""
    try:
        wb = load_workbook(BytesIO(isi), read_only=True, data_only=True)
    except Exception:
        raise FileError("File tidak bisa dibaca. Pastikan file adalah Excel .xlsx yang valid.")

    try:
        if NAMA_SHEET_IMPORT in wb.sheetnames:
            ws = wb[NAMA_SHEET_IMPORT]
        else:
            ws = wb.worksheets[0]

        semua = []
        for nomor, baris in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
            if nomor > MAX_BARIS_SCAN:
                raise FileError(
                    f"File terlalu panjang (lebih dari {MAX_BARIS_SCAN} baris)."
                )
            semua.append((nomor, baris))
    finally:
        wb.close()

    if not semua:
        raise FileError("Sheet kosong.")

    # Baris judul harus di baris 1
    judul = semua[0][1]
    peta = {}
    judul_asli = set()
    for indeks, sel in enumerate(judul):
        nama = _norm(sel)
        if not nama:
            continue
        judul_asli.add(nama)
        kunci = ALIAS_KOLOM.get(nama)
        if kunci and kunci not in peta:
            peta[kunci] = indeks

    if judul_asli & KOLOM_KHAS_EXPORT:
        raise FileError(
            "File ini terlihat seperti hasil Export riwayat (memiliki kolom "
            "No Transaksi/Total). Mengimpornya akan menggandakan transaksi lama. "
            "Gunakan tombol Download Template untuk format import."
        )

    hilang = [
        judul_kolom for judul_kolom, kunci in KOLOM_IMPORT
        if kunci in WAJIB and kunci not in peta
    ]
    if hilang:
        raise FileError(
            "Kolom wajib tidak ditemukan di baris 1: " + ", ".join(hilang)
            + ". Gunakan template dari tombol Download Template."
        )

    return peta, semua[1:]


def _validasi_file(cursor, isi):
    """
    Validasi seluruh file tanpa menyimpan apa pun.

    Return dict:
        errors   : [{"baris": int|None, "pesan": str}]
        warnings : [{"baris": int|None, "pesan": str}]
        rows     : baris valid (siap disimpan)
    """
    peta, baris_data = _baca_sheet(isi)
    _, per_kunci, per_kode = _muat_katalog(cursor)

    errors = []
    warnings = []
    rows = []
    hari_ini = date.today()
    terlihat = {}

    def ambil(baris, kunci):
        indeks = peta.get(kunci)
        if indeks is None or indeks >= len(baris):
            return None
        return baris[indeks]

    for nomor, baris in baris_data:
        kolom_terpakai = [
            ambil(baris, kunci) for _, kunci in KOLOM_IMPORT
        ]
        if all(_teks(v) == "" for v in kolom_terpakai):
            continue   # baris kosong dilewati

        if len(rows) + len(errors) >= MAX_BARIS_DATA * 2:
            raise FileError(f"Terlalu banyak baris data (maksimal {MAX_BARIS_DATA}).")

        galat_baris = []

        # ---- Tanggal
        tanggal = None
        try:
            tanggal = _parse_tanggal(ambil(baris, "tanggal"))
        except ValueError as e:
            galat_baris.append(str(e))

        # ---- Nama penjahit (opsional)
        penjahit = _teks_rapi(ambil(baris, "penjahit"))
        if len(penjahit) > 100:
            galat_baris.append("Nama penjahit maksimal 100 karakter.")

        # ---- Jumlah
        jumlah = None
        try:
            jumlah = _parse_jumlah(ambil(baris, "jumlah"))
        except ValueError as e:
            galat_baris.append(str(e))

        # ---- Harga
        harga = None
        try:
            harga = _parse_harga(ambil(baris, "harga"))
        except ValueError as e:
            galat_baris.append(str(e))

        # ---- Keterangan
        keterangan = _hapus_penanda(_teks(ambil(baris, "keterangan")))

        # ---- Produk & varian (kode produk + warna)
        kode = _teks_rapi(ambil(baris, "kode"))
        warna = _teks_rapi(ambil(baris, "warna"))
        varian = None

        if not kode:
            galat_baris.append("Kode produk wajib diisi.")
        if not warna:
            galat_baris.append("Warna wajib diisi.")

        if kode and warna:
            cocok = per_kunci.get((_norm(kode), _norm(warna)), [])
            if len(cocok) == 1:
                varian = cocok[0]
            elif len(cocok) > 1:
                # Data master ganda (produk + warna sama terdaftar >1 kali).
                # Aman hanya jika tepat satu yang punya baris stok tunggal
                # (satu-satunya yang juga bisa dipakai form manual).
                siap = [v for v in cocok if int(v["jumlah_baris_stok"]) == 1]
                ids = ", ".join(str(v["id_varian"]) for v in cocok)
                if len(siap) == 1:
                    varian = siap[0]
                    warnings.append({
                        "baris": nomor,
                        "pesan": (
                            f"Kode produk '{kode}' + warna '{warna}' terdaftar ganda "
                            f"di database (id_varian {ids}). Dipakai varian "
                            f"{varian['id_varian']} karena hanya itu yang punya data stok."
                        ),
                    })
                else:
                    galat_baris.append(
                        f"Kode produk '{kode}' + warna '{warna}' cocok dengan "
                        f"lebih dari satu varian di database (id_varian {ids}) "
                        f"dan tidak bisa ditentukan otomatis. Rapikan data master dulu."
                    )
            else:
                kandidat = per_kode.get(_norm(kode), [])
                if not kandidat:
                    galat_baris.append(
                        f"Kode produk '{kode}' tidak ditemukan di database."
                    )
                else:
                    tersedia = sorted({v["nama_warna"] for v in kandidat})
                    galat_baris.append(
                        f"Warna '{warna}' tidak tersedia untuk produk '{kode}'. "
                        f"Warna yang ada: {', '.join(tersedia)}."
                    )

        if varian is not None:
            if int(varian["jumlah_baris_stok"]) != 1:
                galat_baris.append(
                    f"Varian {varian['kode_produk']} / {varian['nama_warna']} "
                    f"belum punya data stok & rak yang tunggal "
                    f"(ditemukan {int(varian['jumlah_baris_stok'])} baris stok). "
                    f"Lengkapi data stoknya lebih dulu."
                )

        if galat_baris:
            for pesan in galat_baris:
                errors.append({"baris": nomor, "pesan": pesan})
            continue

        # ---- Peringatan (tidak menggagalkan import)
        nama_info = _teks_rapi(ambil(baris, "nama_produk"))
        if nama_info and _norm(nama_info) != _norm(varian["nama_produk"]):
            warnings.append({
                "baris": nomor,
                "pesan": (
                    f"Nama produk di Excel '{nama_info}' berbeda dengan database "
                    f"'{varian['nama_produk']}'. Sistem memakai kode produk + warna."
                ),
            })
        if tanggal > hari_ini:
            warnings.append({
                "baris": nomor,
                "pesan": f"Tanggal {tanggal.isoformat()} berada di masa depan.",
            })
        if harga == 0:
            warnings.append({"baris": nomor, "pesan": "Harga beli bernilai 0."})
        if varian["status_produk"] == "Non Aktif":
            warnings.append({
                "baris": nomor,
                "pesan": f"Produk {varian['kode_produk']} berstatus Non Aktif.",
            })

        kunci_ganda = (tanggal, _norm(penjahit), varian["id_varian"], jumlah, harga)
        if kunci_ganda in terlihat:
            warnings.append({
                "baris": nomor,
                "pesan": (
                    f"Baris ini identik dengan baris {terlihat[kunci_ganda]} "
                    f"(tanggal, penjahit, varian, jumlah, harga sama). "
                    f"Pastikan bukan input ganda."
                ),
            })
        else:
            terlihat[kunci_ganda] = nomor

        rows.append({
            "baris": nomor,
            "tanggal": tanggal,
            "penjahit": penjahit,
            "id_varian": int(varian["id_varian"]),
            "kode_produk": varian["kode_produk"],
            "nama_produk": varian["nama_produk"],
            "warna": varian["nama_warna"],
            "kode_rak": varian["kode_rak"],
            "jumlah": jumlah,
            "harga": harga,
            "keterangan": keterangan,
        })

    if not rows and not errors:
        raise FileError("Tidak ada baris data di file. Isi data mulai dari baris 2.")
    if len(rows) > MAX_BARIS_DATA:
        raise FileError(f"Terlalu banyak baris data (maksimal {MAX_BARIS_DATA} per file).")

    return {"errors": errors, "warnings": warnings, "rows": rows}


def _hitung_penanda(rows):
    """
    Sidik jari isi file. Dihitung dari data hasil validasi (bukan byte file)
    sehingga file yang disimpan ulang di Excel tetap dikenali sebagai sama,
    dan urutan baris tidak berpengaruh.
    """
    payload = sorted(
        [
            r["tanggal"].isoformat(),
            _norm(r["penjahit"]),
            r["id_varian"],
            r["jumlah"],
            str(r["harga"]),
            _norm(r["keterangan"]),
        ]
        for r in rows
    )
    digest = hashlib.sha256(
        json.dumps(payload, ensure_ascii=False).encode("utf-8")
    ).hexdigest()[:16]
    return f"{PENANDA_AWAL}{digest}]"


def _cari_import_sebelumnya(cursor, penanda):
    cursor.execute(
        """
        SELECT no_faktur, tanggal
        FROM barang_masuk
        WHERE keterangan LIKE %s
        ORDER BY id_barang_masuk
        """,
        (f"%{penanda}%",),
    )
    return cursor.fetchall()


def _nomor_faktur_berikutnya(cursor):
    """Nomor urut BM berikutnya (angka). Format akhir: BM-0001."""
    cursor.execute("""
        SELECT COALESCE(MAX(CAST(SUBSTRING(no_faktur, 4) AS UNSIGNED)), 0)
               AS nomor_terakhir
        FROM barang_masuk
        WHERE no_faktur REGEXP '^BM-[0-9]+$'
    """)
    baris = cursor.fetchone()
    # Kursor biasa mengembalikan tuple, kursor dictionary mengembalikan dict.
    terakhir = baris["nomor_terakhir"] if isinstance(baris, dict) else baris[0]
    return int(terakhir) + 1


def _format_faktur(nomor):
    return f"BM-{nomor:04d}"


def _ringkasan(rows):
    total_pcs = sum(r["jumlah"] for r in rows)
    total_nilai = sum(r["jumlah"] * r["harga"] for r in rows)
    return {
        "jumlah_transaksi": len(rows),
        "jumlah_varian": len({r["id_varian"] for r in rows}),
        "total_pcs": total_pcs,
        "total_nilai": float(total_nilai),
    }


def _respon_error(pesan, status=400):
    return jsonify(ok=False, errors=[{"baris": None, "pesan": pesan}]), status


# =========================================================
# PREVIEW IMPORT (tidak menyimpan apa pun)
# =========================================================

@barang_masuk_excel_bp.route("/barang-masuk/import/preview", methods=["POST"])
def import_preview():
    isi, galat = _ambil_file_upload()
    if galat:
        return _respon_error(galat)

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        try:
            hasil = _validasi_file(cursor, isi)
        except FileError as e:
            return _respon_error(str(e))

        rows = hasil["rows"]
        respon = {
            "ok": not hasil["errors"],
            "errors": hasil["errors"],
            "warnings": hasil["warnings"],
            "ringkasan": _ringkasan(rows),
            "jumlah_baris_valid": len(rows),
            "pratinjau": [],
            "duplikat": None,
            "faktur_perkiraan": None,
        }

        if not hasil["errors"]:
            respon["pratinjau"] = [
                {
                    "baris": r["baris"],
                    "tanggal": r["tanggal"].isoformat(),
                    "penjahit": r["penjahit"],
                    "kode_produk": r["kode_produk"],
                    "nama_produk": r["nama_produk"],
                    "warna": r["warna"],
                    "kode_rak": r["kode_rak"],
                    "jumlah": r["jumlah"],
                    "harga": float(r["harga"]),
                    "total": float(r["jumlah"] * r["harga"]),
                }
                for r in rows[:MAX_PREVIEW_BARIS]
            ]
            respon["pratinjau_dipotong"] = len(rows) > MAX_PREVIEW_BARIS

            awal = _nomor_faktur_berikutnya(cursor)
            respon["faktur_perkiraan"] = (
                f"{_format_faktur(awal)} s/d {_format_faktur(awal + len(rows) - 1)}"
            )

            sebelumnya = _cari_import_sebelumnya(cursor, _hitung_penanda(rows))
            if sebelumnya:
                respon["duplikat"] = {
                    "jumlah": len(sebelumnya),
                    "no_faktur": [s["no_faktur"] for s in sebelumnya][:10],
                }

        return jsonify(respon)

    except Exception as e:
        print("ERROR PREVIEW IMPORT BARANG MASUK:", e)
        return _respon_error("Terjadi kesalahan saat memeriksa file.", 500)
    finally:
        cursor.close()
        conn.close()


# =========================================================
# KONFIRMASI IMPORT (simpan transaksi + tambah stok)
# =========================================================

@barang_masuk_excel_bp.route("/barang-masuk/import/konfirmasi", methods=["POST"])
def import_konfirmasi():
    isi, galat = _ambil_file_upload()
    if galat:
        return _respon_error(galat)

    izinkan_duplikat = request.form.get("izinkan_duplikat") == "1"

    conn = get_db_connection()
    cursor = None
    try:
        # ---- Tahap 1: validasi ulang penuh (data bisa berubah sejak pratinjau)
        baca = conn.cursor(dictionary=True)
        try:
            hasil = _validasi_file(baca, isi)
        except FileError as e:
            return _respon_error(str(e))
        finally:
            baca.close()

        if hasil["errors"]:
            return jsonify(
                ok=False,
                errors=hasil["errors"],
            ), 400

        rows = hasil["rows"]
        penanda = _hitung_penanda(rows)

        # Akhiri transaksi baca agar tahap tulis memakai data terbaru
        conn.rollback()

        # ---- Tahap 2: tulis dalam SATU transaksi
        cursor = conn.cursor()

        # Kunci stok semua varian (urut id_varian agar tidak deadlock)
        for id_varian in sorted({r["id_varian"] for r in rows}):
            cursor.execute(
                "SELECT id_stok FROM stok_barang WHERE id_varian = %s FOR UPDATE",
                (id_varian,),
            )
            if len(cursor.fetchall()) != 1:
                raise ValueError(
                    "Data stok salah satu varian tidak ditemukan atau ganda."
                )

        # Cek file sama setelah stok terkunci, agar dua import bersamaan
        # dari file yang sama tidak lolos keduanya.
        if not izinkan_duplikat:
            cursor_cek = conn.cursor(dictionary=True)
            try:
                sebelumnya = _cari_import_sebelumnya(cursor_cek, penanda)
            finally:
                cursor_cek.close()
            if sebelumnya:
                conn.rollback()
                return jsonify(
                    ok=False,
                    duplikat=True,
                    errors=[{
                        "baris": None,
                        "pesan": (
                            "File dengan isi yang sama sudah pernah diimpor "
                            f"({', '.join(s['no_faktur'] for s in sebelumnya[:5])}"
                            f"{', ...' if len(sebelumnya) > 5 else ''}). "
                            "Centang konfirmasi proses ulang jika memang disengaja."
                        ),
                    }],
                ), 409

        nomor = _nomor_faktur_berikutnya(cursor)
        faktur_awal = _format_faktur(nomor)

        for r in rows:
            no_faktur = _format_faktur(nomor)
            nomor += 1
            keterangan = f"{r['keterangan']} {penanda}".strip()

            cursor.execute(
                """
                INSERT INTO barang_masuk
                    (no_faktur, nama_penjahit, tanggal, keterangan)
                VALUES (%s, %s, %s, %s)
                """,
                (no_faktur, r["penjahit"], r["tanggal"], keterangan),
            )
            id_barang_masuk = cursor.lastrowid

            cursor.execute(
                """
                INSERT INTO detail_barang_masuk
                    (id_barang_masuk, id_varian, jumlah, harga_beli)
                VALUES (%s, %s, %s, %s)
                """,
                (id_barang_masuk, r["id_varian"], r["jumlah"], r["harga"]),
            )

            cursor.execute(
                "UPDATE stok_barang SET stok = stok + %s WHERE id_varian = %s",
                (r["jumlah"], r["id_varian"]),
            )
            if cursor.rowcount != 1:
                raise ValueError(
                    f"Stok varian {r['kode_produk']} / {r['warna']} "
                    f"gagal diperbarui (baris Excel {r['baris']})."
                )

        faktur_akhir = _format_faktur(nomor - 1)
        conn.commit()

        ringkas = _ringkasan(rows)
        return jsonify(
            ok=True,
            faktur_awal=faktur_awal,
            faktur_akhir=faktur_akhir,
            **ringkas,
        )

    except Exception as e:
        conn.rollback()
        print("ERROR IMPORT BARANG MASUK:", e)
        if isinstance(e, ValueError):
            pesan = str(e)
        else:
            pesan = "Terjadi kesalahan pada database."
        return _respon_error(
            f"Import dibatalkan, tidak ada data yang tersimpan. {pesan}",
            500,
        )
    finally:
        if cursor is not None:
            cursor.close()
        conn.close()


# =========================================================
# STYLE EXCEL
# =========================================================

_TIPIS = Side(style="thin", color="E0CDD0")
_BORDER = Border(left=_TIPIS, right=_TIPIS, top=_TIPIS, bottom=_TIPIS)


def _gaya_header(ws, baris, jumlah_kolom):
    for kolom in range(1, jumlah_kolom + 1):
        sel = ws.cell(row=baris, column=kolom)
        sel.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        sel.fill = PatternFill("solid", start_color=WARNA_HEADER)
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sel.border = _BORDER
    ws.row_dimensions[baris].height = 28


def _atur_lebar(ws, lebar):
    for indeks, nilai in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(indeks)].width = nilai


def _kirim_workbook(wb, nama_file):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=nama_file,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# =========================================================
# EXPORT RIWAYAT BARANG MASUK
# =========================================================

@barang_masuk_excel_bp.route("/barang-masuk/export-excel")
def export_barang_masuk_excel():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        # Hanya membaca data; tidak mengubah stok maupun transaksi.
        cursor.execute("""
            SELECT bm.no_faktur,
                   bm.tanggal,
                   bm.nama_penjahit,
                   p.kode_produk,
                   p.nama_produk,
                   w.nama_warna,
                   (SELECT MIN(r.kode_rak)
                      FROM stok_barang sb
                      JOIN rak r ON r.id_rak = sb.id_rak
                     WHERE sb.id_varian = vp.id_varian) AS kode_rak,
                   dbm.jumlah,
                   dbm.harga_beli,
                   bm.keterangan
            FROM barang_masuk bm
            JOIN detail_barang_masuk dbm
                 ON dbm.id_barang_masuk = bm.id_barang_masuk
            JOIN varian_produk vp ON vp.id_varian = dbm.id_varian
            JOIN produk p ON p.id_produk = vp.id_produk
            JOIN warna w ON w.id_warna = vp.id_warna
            ORDER BY bm.tanggal DESC, bm.id_barang_masuk DESC, dbm.id_detail
        """)
        data = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = NAMA_SHEET_EXPORT

    judul = [
        "No Transaksi", "Tanggal", "Nama Penjahit", "Kode Produk",
        "Nama Produk", "Warna", "Kode Rak", "Jumlah",
        "Harga Beli per Pcs", "Total", "Keterangan",
    ]
    ws.append(judul)
    _gaya_header(ws, 1, len(judul))

    for nomor, d in enumerate(data, start=2):
        jumlah = int(d["jumlah"])
        harga = float(d["harga_beli"])
        ws.append([
            None, d["tanggal"], None, None, None, None, None,
            jumlah, harga, jumlah * harga, None,
        ])
        _set_teks(ws.cell(row=nomor, column=1), d["no_faktur"])
        _set_teks(ws.cell(row=nomor, column=3), d["nama_penjahit"] or "")
        _set_teks(ws.cell(row=nomor, column=4), d["kode_produk"])
        _set_teks(ws.cell(row=nomor, column=5), d["nama_produk"])
        _set_teks(ws.cell(row=nomor, column=6), d["nama_warna"])
        _set_teks(ws.cell(row=nomor, column=7), d["kode_rak"] or "")
        _set_teks(ws.cell(row=nomor, column=11), _hapus_penanda(d["keterangan"]))

        ws.cell(row=nomor, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=nomor, column=8).number_format = "#,##0"
        ws.cell(row=nomor, column=9).number_format = "#,##0"
        ws.cell(row=nomor, column=10).number_format = "#,##0"
        for kolom in range(1, len(judul) + 1):
            sel = ws.cell(row=nomor, column=kolom)
            sel.font = Font(name=FONT, size=10)
            sel.border = _BORDER
            if kolom in (2, 7):
                sel.alignment = Alignment(horizontal="center")

    _atur_lebar(ws, [14, 13, 18, 13, 26, 16, 10, 9, 18, 16, 34])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(judul))}{max(ws.max_row, 1)}"

    nama_file = f"riwayat_barang_masuk_{date.today().strftime('%Y%m%d')}.xlsx"
    return _kirim_workbook(wb, nama_file)


# =========================================================
# TEMPLATE IMPORT
# =========================================================

@barang_masuk_excel_bp.route("/barang-masuk/template-excel")
def template_barang_masuk_excel():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        daftar_varian, _, _ = _muat_katalog(cursor)
    finally:
        cursor.close()
        conn.close()

    wb = Workbook()

    # ---------------- Sheet 1: tempat mengisi data ----------------
    ws = wb.active
    ws.title = NAMA_SHEET_IMPORT
    judul = [j for j, _ in KOLOM_IMPORT]
    ws.append(judul)
    _gaya_header(ws, 1, len(judul))
    _atur_lebar(ws, [14, 18, 14, 28, 18, 10, 20, 34])
    ws.freeze_panes = "A2"

    for baris in range(2, MAX_BARIS_DATA + 2):
        for kolom in range(1, len(judul) + 1):
            sel = ws.cell(row=baris, column=kolom)
            sel.font = Font(name=FONT, size=10)
        ws.cell(row=baris, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=baris, column=3).number_format = "@"
        ws.cell(row=baris, column=6).number_format = "0"
        ws.cell(row=baris, column=7).number_format = "#,##0"

    batas = MAX_BARIS_DATA + 1
    dv_tanggal = DataValidation(
        type="date", operator="greaterThan", formula1="36526",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Tanggal tidak valid",
        error="Isi dengan tanggal yang valid, contoh 2026-10-20.",
    )
    dv_jumlah = DataValidation(
        type="whole", operator="greaterThan", formula1="0",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Jumlah tidak valid",
        error="Jumlah harus bilangan bulat lebih dari 0.",
    )
    dv_harga = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1="0",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Harga tidak valid",
        error="Harga beli harus angka, tanpa titik/koma pemisah ribuan.",
    )
    ws.add_data_validation(dv_tanggal)
    ws.add_data_validation(dv_jumlah)
    ws.add_data_validation(dv_harga)
    dv_tanggal.add(f"A2:A{batas}")
    dv_jumlah.add(f"F2:F{batas}")
    dv_harga.add(f"G2:G{batas}")

    # ---------------- Sheet 2: petunjuk + contoh ----------------
    wp = wb.create_sheet(NAMA_SHEET_PETUNJUK)
    contoh = daftar_varian[0] if daftar_varian else None
    contoh_kode = contoh["kode_produk"] if contoh else "MK001"
    contoh_nama = contoh["nama_produk"] if contoh else "Mukena Abaya"
    contoh_warna = contoh["nama_warna"] if contoh else "Black"

    petunjuk = [
        ("PETUNJUK IMPORT BARANG MASUK", "judul"),
        ("", None),
        ("1. Isi data pada sheet 'Import Barang Masuk', mulai dari baris 2. Jangan mengubah judul kolom di baris 1.", None),
        ("2. Satu baris = satu varian (produk + warna) yang diterima. Setiap baris akan menjadi satu transaksi BM-xxxx.", None),
        ("3. Kolom WAJIB: Tanggal, Kode Produk, Warna, Jumlah, Harga Beli per Pcs.", None),
        ("4. Kolom opsional: Nama Penjahit, Nama Produk (info saja), Keterangan.", None),
        ("5. Barang dikenali dari KODE PRODUK + WARNA. Salin persis dari sheet 'Daftar Varian'. Nama produk tidak dipakai untuk mencocokkan.", None),
        ("6. Tanggal memakai format YYYY-MM-DD (contoh 2026-10-20). Jumlah bilangan bulat > 0. Harga beli angka biasa tanpa titik/koma (contoh 150000).", None),
        ("7. Rak tidak perlu diisi: rak mengikuti data stok varian yang sudah ada.", None),
        ("8. Setelah upload, sistem menampilkan pratinjau. Stok baru bertambah setelah Anda menekan Konfirmasi.", None),
        ("9. File dengan isi yang sama tidak akan diproses dua kali kecuali Anda mencentang konfirmasi proses ulang.", None),
        ("10. Jangan meng-import file hasil Export riwayat; formatnya berbeda dan akan menggandakan transaksi lama.", None),
        ("", None),
        ("CONTOH PENGISIAN (jangan disalin sebagai data kecuali memang diterima):", "sub"),
    ]
    for teks, gaya in petunjuk:
        wp.append([teks])
        sel = wp.cell(row=wp.max_row, column=1)
        if gaya == "judul":
            sel.font = Font(name=FONT, bold=True, size=14, color="B5606C")
        elif gaya == "sub":
            sel.font = Font(name=FONT, bold=True, size=10)
        else:
            sel.font = Font(name=FONT, size=10)

    baris_judul = wp.max_row + 1
    wp.append(judul)
    _gaya_header(wp, baris_judul, len(judul))
    wp.append([
        "2026-10-20", "Alyya", contoh_kode, contoh_nama, contoh_warna,
        12, 150000, "Kiriman pertama bulan ini",
    ])
    for kolom in range(1, len(judul) + 1):
        wp.cell(row=wp.max_row, column=kolom).font = Font(name=FONT, size=10)
        wp.cell(row=wp.max_row, column=kolom).border = _BORDER
    _atur_lebar(wp, [14, 18, 14, 28, 18, 10, 20, 34])

    # ---------------- Sheet 3: daftar varian ----------------
    wv = wb.create_sheet(NAMA_SHEET_VARIAN)
    judul_varian = ["Kode Produk", "Nama Produk", "Warna", "Kode Rak", "Stok Saat Ini", "Status Produk"]
    wv.append(judul_varian)
    _gaya_header(wv, 1, len(judul_varian))
    for nomor, v in enumerate(daftar_varian, start=2):
        wv.append([None, None, None, None, int(v["stok"]), None])
        _set_teks(wv.cell(row=nomor, column=1), v["kode_produk"])
        _set_teks(wv.cell(row=nomor, column=2), v["nama_produk"])
        _set_teks(wv.cell(row=nomor, column=3), v["nama_warna"])
        _set_teks(wv.cell(row=nomor, column=4), v["kode_rak"] or "")
        _set_teks(wv.cell(row=nomor, column=6), v["status_produk"] or "")
        for kolom in range(1, len(judul_varian) + 1):
            sel = wv.cell(row=nomor, column=kolom)
            sel.font = Font(name=FONT, size=10)
            sel.border = _BORDER
    _atur_lebar(wv, [14, 28, 18, 10, 14, 14])
    wv.freeze_panes = "A2"
    wv.auto_filter.ref = f"A1:F{max(wv.max_row, 1)}"

    return _kirim_workbook(wb, "template_import_barang_masuk.xlsx")