from flask import Flask, render_template, request, redirect, flash, send_file
from models.database import get_db_connection
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import os
import uuid
from datetime import datetime
from werkzeug.utils import secure_filename

app = Flask(__name__)


# =========================================================
# HOME
# =========================================================

@app.route("/")
def home():
    return "Backend YN Hijab berhasil dijalankan!"


# =========================================================
# TEST DATABASE
# =========================================================

@app.route("/test-db")
def test_db():
    try:
        conn = get_db_connection()
        conn.close()
        return "Koneksi database berhasil!"
    except Exception as e:
        return f"Koneksi database gagal: {e}"


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/dashboard")
def dashboard():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        # Total produk aktif
        cursor.execute("""
            SELECT COUNT(*) AS total_produk
            FROM produk
            WHERE status = 'Aktif'
        """)
        total_produk = cursor.fetchone()["total_produk"]

        # Total stok
        cursor.execute("""
            SELECT COALESCE(SUM(stok), 0) AS total_stok
            FROM stok_barang
        """)
        total_stok = cursor.fetchone()["total_stok"]

        # Pesanan hari ini
        cursor.execute("""
            SELECT COUNT(*) AS pesanan_hari_ini
            FROM pesanan
            WHERE DATE(tanggal) = CURDATE()
              AND STATUS != 'dibatalkan'
        """)
        pesanan_hari_ini = cursor.fetchone()["pesanan_hari_ini"]

        # Stok menipis
        cursor.execute("""
            SELECT COUNT(*) AS stok_menipis
            FROM stok_barang
            WHERE stok <= stok_minimum
        """)
        stok_menipis = cursor.fetchone()["stok_menipis"]

        # Penjualan 7 hari
        cursor.execute("""
            SELECT DATE_FORMAT(tanggal, '%Y-%m-%d') AS tanggal,
                   CAST(COALESCE(SUM(total_harga), 0) AS DOUBLE) AS penjualan
            FROM pesanan
            WHERE STATUS != 'dibatalkan'
              AND tanggal >= CURDATE() - INTERVAL 6 DAY
              AND tanggal < CURDATE() + INTERVAL 1 DAY
            GROUP BY DATE(tanggal)
            ORDER BY tanggal ASC
        """)
        data_7_hari = cursor.fetchall()

        # Penjualan 30 hari
        cursor.execute("""
            SELECT DATE_FORMAT(tanggal, '%Y-%m-%d') AS tanggal,
                   CAST(COALESCE(SUM(total_harga), 0) AS DOUBLE) AS penjualan
            FROM pesanan
            WHERE STATUS != 'dibatalkan'
              AND tanggal >= CURDATE() - INTERVAL 29 DAY
              AND tanggal < CURDATE() + INTERVAL 1 DAY
            GROUP BY DATE(tanggal)
            ORDER BY tanggal ASC
        """)
        data_30_hari = cursor.fetchall()

        # Penjualan 3 bulan
        cursor.execute("""
            SELECT DATE_FORMAT(tanggal, '%Y-%m') AS bulan,
                   CAST(COALESCE(SUM(total_harga), 0) AS DOUBLE) AS penjualan
            FROM pesanan
            WHERE STATUS != 'dibatalkan'
              AND tanggal >= DATE_FORMAT(
                    CURDATE() - INTERVAL 2 MONTH,
                    '%Y-%m-01'
                  )
              AND tanggal < DATE_FORMAT(
                    CURDATE() + INTERVAL 1 MONTH,
                    '%Y-%m-01'
                  )
            GROUP BY DATE_FORMAT(tanggal, '%Y-%m')
            ORDER BY bulan ASC
        """)
        data_3_bulan = cursor.fetchall()

        # Penjualan 1 tahun
        cursor.execute("""
            SELECT DATE_FORMAT(tanggal, '%Y-%m') AS bulan,
                   CAST(COALESCE(SUM(total_harga), 0) AS DOUBLE) AS penjualan
            FROM pesanan
            WHERE STATUS != 'dibatalkan'
              AND tanggal >= DATE_FORMAT(
                    CURDATE() - INTERVAL 11 MONTH,
                    '%Y-%m-01'
                  )
              AND tanggal < DATE_FORMAT(
                    CURDATE() + INTERVAL 1 MONTH,
                    '%Y-%m-01'
                  )
            GROUP BY DATE_FORMAT(tanggal, '%Y-%m')
            ORDER BY bulan ASC
        """)
        data_1_tahun = cursor.fetchall()

        # Stok menipis
        cursor.execute("""
            SELECT p.nama_produk,
                   w.nama_warna,
                   sb.stok,
                   sb.stok_minimum
            FROM stok_barang sb
            JOIN varian_produk vp
                ON sb.id_varian = vp.id_varian
            JOIN produk p
                ON vp.id_produk = p.id_produk
            JOIN warna w
                ON vp.id_warna = w.id_warna
            WHERE p.status = 'Aktif'
              AND sb.stok <= sb.stok_minimum
            ORDER BY sb.stok ASC
            LIMIT 5
        """)
        data_stok_menipis = cursor.fetchall()

        # Pesanan terbaru
        cursor.execute("""
            SELECT id_pesanan,
                   no_pesanan,
                   DATE_FORMAT(tanggal, '%Y-%m-%d %H:%i') AS tanggal,
                   nama_customor,
                   CAST(
                       COALESCE(total_harga, 0)
                       AS DOUBLE
                   ) AS total_harga,
                   STATUS
            FROM pesanan
            WHERE STATUS != 'dibatalkan'
            ORDER BY tanggal DESC
            LIMIT 3
        """)
        pesanan_terbaru = cursor.fetchall()

        return render_template(
            "dashboard.html",
            total_produk=total_produk,
            total_stok=total_stok,
            pesanan_hari_ini=pesanan_hari_ini,
            stok_menipis=stok_menipis,
            data_7_hari=data_7_hari,
            data_30_hari=data_30_hari,
            data_3_bulan=data_3_bulan,
            data_1_tahun=data_1_tahun,
            data_stok_menipis=data_stok_menipis,
            pesanan_terbaru=pesanan_terbaru
        )

    finally:
        cursor.close()
        conn.close()


# =========================================================
# PRODUK
# =========================================================

@app.route("/produk")
def produk():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT COUNT(*) AS total_produk
            FROM produk
            WHERE status = 'Aktif'
        """)
        total_produk = cursor.fetchone()["total_produk"]

        cursor.execute("""
            SELECT COALESCE(SUM(stok), 0) AS total_stok
            FROM stok_barang
        """)
        total_stok = cursor.fetchone()["total_stok"]

        cursor.execute("""
            SELECT COUNT(*) AS stok_menipis
            FROM stok_barang
            WHERE stok <= stok_minimum
        """)
        stok_menipis = cursor.fetchone()["stok_menipis"]

        cursor.execute("""
            SELECT p.id_produk,
                   p.kode_produk,
                   p.nama_produk,
                   p.harga_jual,
                   p.status,
                   p.gambar,
                   k.nama_kategori,
                   w.nama_warna,
                   COALESCE(sb.stok, 0) AS stok,
                   r.kode_rak AS rak
            FROM produk p
            LEFT JOIN kategori k
                ON p.id_kategori = k.id_kategori
            LEFT JOIN varian_produk vp
                ON p.id_produk = vp.id_produk
            LEFT JOIN warna w
                ON vp.id_warna = w.id_warna
            LEFT JOIN stok_barang sb
                ON vp.id_varian = sb.id_varian
            LEFT JOIN rak r
                ON sb.id_rak = r.id_rak
            WHERE p.status = 'Aktif'
            ORDER BY p.id_produk DESC
        """)
        data_produk = cursor.fetchall()

        cursor.execute("""
            SELECT id_kategori,
                   nama_kategori
            FROM kategori
            ORDER BY nama_kategori
        """)
        data_kategori = cursor.fetchall()

        cursor.execute("""
            SELECT id_warna,
                   nama_warna
            FROM warna
            ORDER BY nama_warna
        """)
        data_warna = cursor.fetchall()

        return render_template(
            "produk.html",
            produk=data_produk,
            kategori=data_kategori,
            warna=data_warna,
            total_produk=total_produk,
            total_stok=total_stok,
            stok_menipis=stok_menipis
        )

    finally:
        cursor.close()
        conn.close()


@app.route("/produk/tambah", methods=["POST"])
def tambah_produk():
    nama_produk = request.form["nama_produk"]
    id_kategori = request.form["id_kategori"]
    harga_jual = request.form["harga_jual"].replace(".", "")
    id_rak = request.form["id_rak"]
    stok_awal = request.form.get("stok_awal", 0)
    warna_dipilih = request.form.getlist("warna[]")

    gambar = request.files.get("gambar")
    nama_file_gambar = None

    if gambar and gambar.filename:
        nama_file_gambar = secure_filename(gambar.filename)

        folder_gambar = os.path.join(
            app.root_path,
            "static",
            "img",
            "produk"
        )

        os.makedirs(folder_gambar, exist_ok=True)

        gambar.save(
            os.path.join(
                folder_gambar,
                nama_file_gambar
            )
        )

    conn = get_db_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT kode_produk
            FROM produk
            WHERE kode_produk LIKE 'PR%'
            ORDER BY id_produk DESC
            LIMIT 1
        """)

        data_terakhir = cursor.fetchone()

        if data_terakhir:
            nomor_lama = int(
                data_terakhir[0].replace("PR", "")
            )
            nomor_baru = nomor_lama + 1
        else:
            nomor_baru = 1

        kode_produk = f"PR{nomor_baru:03d}"

        cursor.execute("""
            INSERT INTO produk
            (
                id_kategori,
                kode_produk,
                nama_produk,
                harga_jual,
                gambar,
                status
            )
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (
            id_kategori,
            kode_produk,
            nama_produk,
            harga_jual,
            nama_file_gambar,
            "Aktif"
        ))

        id_produk = cursor.lastrowid

        for id_warna in warna_dipilih:

            cursor.execute("""
                INSERT INTO varian_produk
                (id_produk, id_warna)
                VALUES (%s, %s)
            """, (
                id_produk,
                id_warna
            ))

            id_varian = cursor.lastrowid

            cursor.execute("""
                INSERT INTO stok_barang
                (
                    id_varian,
                    id_rak,
                    stok,
                    stok_minimum
                )
                VALUES (%s, %s, %s, %s)
            """, (
                id_varian,
                id_rak,
                stok_awal,
                5
            ))

        conn.commit()

    except Exception as e:
        conn.rollback()
        print("ERROR TAMBAH PRODUK:", e)
        raise

    finally:
        cursor.close()
        conn.close()

    return redirect("/produk")


@app.route("/produk/hapus/<int:id_produk>", methods=["POST"])
def hapus_produk(id_produk):
    conn = get_db_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            UPDATE produk
            SET status = 'Nonaktif'
            WHERE id_produk = %s
        """, (id_produk,))

        conn.commit()

    except Exception:
        conn.rollback()
        raise

    finally:
        cursor.close()
        conn.close()

    return redirect("/produk")


# =========================================================
# NOMOR PESANAN
# =========================================================

def generate_no_pesanan(cursor, tanggal):
    # Format:
    # #0001-DDMMYY
    #
    # Nomor urut reset setiap 6 bulan:
    # Januari - Juni
    # Juli - Desember

    dt = datetime.fromisoformat(tanggal)

    if dt.month <= 6:

        mulai = dt.replace(
            month=1,
            day=1,
            hour=0,
            minute=0,
            second=0,
            microsecond=0
        )

        selesai = dt.replace(
            month=7,
            day=1,
            hour=0,
            minute=0,
            second=0,
            microsecond=0
        )

    else:

        mulai = dt.replace(
            month=7,
            day=1,
            hour=0,
            minute=0,
            second=0,
            microsecond=0
        )

        selesai = datetime(
            dt.year + 1,
            1,
            1
        )

    cursor.execute("""
        SELECT COALESCE(
            MAX(
                CAST(
                    SUBSTRING(no_pesanan, 2, 4)
                    AS UNSIGNED
                )
            ),
            0
        ) AS nomor_terakhir
        FROM pesanan
        WHERE no_pesanan IS NOT NULL
          AND tanggal >= %s
          AND tanggal < %s
    """, (
        mulai,
        selesai
    ))

    data_nomor = cursor.fetchone()

    nomor_baru = int(
        data_nomor["nomor_terakhir"] or 0
    ) + 1

    if nomor_baru > 9999:
        raise Exception(
            "Nomor pesanan periode ini sudah mencapai 9999."
        )

    return (
        f"#{nomor_baru:04d}-"
        f"{dt.strftime('%d%m%y')}"
    )


# =========================================================
# PESANAN
# =========================================================

@app.route("/pesanan")
def pesanan():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT *
            FROM pesanan
            ORDER BY id_pesanan DESC
        """)

        pesanan_data = cursor.fetchall()

        cursor.execute("""
            SELECT COUNT(*) AS total_pesanan,

                   SUM(
                       CASE
                           WHEN STATUS = 'Menunggu'
                           THEN 1
                           ELSE 0
                       END
                   ) AS menunggu,

                   COUNT(
                       CASE
                           WHEN STATUS = 'Selesai'
                           THEN 1
                       END
                   ) AS selesai,

                   COALESCE(
                       SUM(total_harga),
                       0
                   ) AS total_penjualan

            FROM pesanan
        """)

        ringkasan = cursor.fetchone()

        cursor.execute("""
            SELECT vp.id_varian,
                   p.nama_produk,
                   w.nama_warna,
                   p.harga_jual,
                   COALESCE(sb.stok, 0) AS stok

            FROM varian_produk vp

            JOIN produk p
                ON vp.id_produk = p.id_produk

            JOIN warna w
                ON vp.id_warna = w.id_warna

            LEFT JOIN stok_barang sb
                ON vp.id_varian = sb.id_varian

            WHERE p.status = 'Aktif'

            ORDER BY p.nama_produk,
                     w.nama_warna
        """)

        varian = cursor.fetchall()

        return render_template(
            "pesanan.html",
            pesanan=pesanan_data,
            varian=varian,
            ringkasan=ringkasan
        )

    finally:
        cursor.close()
        conn.close()


@app.route("/pesanan/detail/<int:id_pesanan>")
def detail_pesanan(id_pesanan):
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT id_pesanan,
                   no_pesanan,
                   no_hp,
                   nama_customor,
                   alamat,
                   tanggal,

                   DATE_FORMAT(
                       tanggal,
                       '%Y-%m-%dT%H:%i'
                   ) AS tanggal_edit,

                   STATUS,
                   total_harga,
                   keterangan,
                   metode_pengambilan

            FROM pesanan

            WHERE id_pesanan = %s
        """, (id_pesanan,))

        pesanan_data = cursor.fetchone()

        if not pesanan_data:
            return {
                "error": "Pesanan tidak ditemukan"
            }, 404

        cursor.execute("""
            SELECT dp.id_detail,
                   dp.id_varian,
                   p.nama_produk,
                   w.nama_warna,
                   dp.harga_jual,
                   dp.jumlah_pesanan,
                   dp.subtotal

            FROM detail_pesanan dp

            JOIN varian_produk vp
                ON dp.id_varian = vp.id_varian

            JOIN produk p
                ON vp.id_produk = p.id_produk

            JOIN warna w
                ON vp.id_warna = w.id_warna

            WHERE dp.id_pesanan = %s

            ORDER BY dp.id_detail ASC
        """, (id_pesanan,))

        detail = cursor.fetchall()

        return {
            "pesanan": pesanan_data,
            "detail": detail
        }

    except Exception as e:
        print("ERROR DETAIL PESANAN:", e)

        return {
            "error": str(e)
        }, 500

    finally:
        cursor.close()
        conn.close()


@app.route("/pesanan/varian-edit")
def varian_edit():

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        cursor.execute("""
            SELECT
                vp.id_varian,
                p.nama_produk,
                w.nama_warna,
                p.harga_jual,
                COALESCE(sb.stok, 0) AS stok
            FROM varian_produk vp
            JOIN produk p
                ON vp.id_produk = p.id_produk
            JOIN warna w
                ON vp.id_warna = w.id_warna
            LEFT JOIN stok_barang sb
                ON vp.id_varian = sb.id_varian
            WHERE p.status = 'Aktif'
            ORDER BY p.nama_produk, w.nama_warna
        """)

        data = cursor.fetchall()

        return data

    finally:

        cursor.close()
        conn.close()

# =========================================================
# TAMBAH PESANAN
# =========================================================

@app.route("/pesanan/tambah", methods=["POST"])
def tambah_pesanan():

    no_hp = request.form["no_hp"]
    nama_customer = request.form["nama_customer"]
    alamat = request.form.get("alamat", "")
    tanggal = request.form["tanggal"]
    status = request.form["status"]
    keterangan = request.form.get("keterangan", "")
    metode_pengambilan = request.form["metode_pengambilan"]

    id_varian_list = request.form.getlist(
        "id_varian[]"
    )

    jumlah_list = request.form.getlist(
        "jumlah[]"
    )

    if not id_varian_list:
        return (
            "Minimal satu barang harus dipilih.",
            400
        )

    if len(id_varian_list) != len(jumlah_list):
        return (
            "Data barang pesanan tidak valid.",
            400
        )

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        no_pesanan = generate_no_pesanan(
            cursor,
            tanggal
        )

        cursor.execute("""
            INSERT INTO pesanan
            (
                no_pesanan,
                no_hp,
                nama_customor,
                alamat,
                tanggal,
                status,
                total_harga,
                keterangan,
                metode_pengambilan
            )

            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )
        """, (
            no_pesanan,
            no_hp,
            nama_customer,
            alamat,
            tanggal,
            status,
            0,
            keterangan,
            metode_pengambilan
        ))

        id_pesanan = cursor.lastrowid
        total_harga = 0

        varian_dipakai = set()

        for id_varian_text, jumlah_text in zip(
            id_varian_list,
            jumlah_list
        ):

            id_varian = int(id_varian_text)
            jumlah = int(jumlah_text)

            if id_varian in varian_dipakai:
                raise Exception(
                    "Varian produk yang sama tidak boleh "
                    "dipilih dua kali."
                )

            varian_dipakai.add(id_varian)

            if jumlah <= 0:
                raise Exception(
                    "Jumlah pesanan harus lebih dari 0."
                )

            cursor.execute("""
                SELECT p.harga_jual,
                       sb.stok

                FROM varian_produk vp

                JOIN produk p
                    ON vp.id_produk = p.id_produk

                JOIN stok_barang sb
                    ON vp.id_varian = sb.id_varian

                WHERE vp.id_varian = %s

                FOR UPDATE
            """, (id_varian,))

            data = cursor.fetchone()

            if not data:
                raise Exception(
                    "Stok varian tidak ditemukan."
                )

            harga = float(
                data["harga_jual"] or 0
            )

            stok = int(
                data["stok"] or 0
            )

            if jumlah > stok:
                raise Exception(
                    f"Stok tidak cukup. "
                    f"Stok tersedia: {stok}, "
                    f"jumlah yang diminta: {jumlah}."
                )

            subtotal = harga * jumlah

            total_harga += subtotal

            cursor.execute("""
                INSERT INTO detail_pesanan
                (
                    id_pesanan,
                    id_varian,
                    jumlah_pesanan,
                    harga_jual,
                    subtotal
                )

                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                id_pesanan,
                id_varian,
                jumlah,
                harga,
                subtotal
            ))

            cursor.execute("""
                INSERT INTO barang_keluar
                (
                    id_varian,
                    jumlah,
                    alasan,
                    tanggal,
                    keterangan
                )

                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                id_varian,
                jumlah,
                "Penjualan",
                tanggal.split("T")[0],
                f"Pesanan {no_pesanan}"
            ))

            cursor.execute("""
                UPDATE stok_barang
                SET stok = stok - %s

                WHERE id_varian = %s
                  AND stok >= %s
            """, (
                jumlah,
                id_varian,
                jumlah
            ))

            if cursor.rowcount != 1:
                raise Exception(
                    "Stok tidak mencukupi."
                )

        cursor.execute("""
            UPDATE pesanan
            SET total_harga = %s
            WHERE id_pesanan = %s
        """, (
            total_harga,
            id_pesanan
        ))

        conn.commit()

    except Exception as e:

        conn.rollback()

        print("ERROR PESANAN:", e)

        return (
            f"Pesanan gagal disimpan: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/pesanan")


# =========================================================
# EDIT PESANAN
# =========================================================

@app.route("/pesanan/edit/<int:id_pesanan>", methods=["POST"])
def update_pesanan(id_pesanan):

    no_hp = request.form["no_hp"]
    nama_customer = request.form["nama_customer"]
    alamat = request.form.get("alamat", "")
    tanggal = request.form["tanggal"]
    status = request.form["status"]
    metode_pengambilan = request.form["metode_pengambilan"]
    keterangan = request.form.get("keterangan", "")

    id_varian_list = request.form.getlist(
        "id_varian[]"
    )

    jumlah_list = request.form.getlist(
        "jumlah[]"
    )

    if not id_varian_list:
        return (
            "Minimal satu barang harus dipilih.",
            400
        )

    if len(id_varian_list) != len(jumlah_list):
        return (
            "Data barang pesanan tidak valid.",
            400
        )

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        # -----------------------------------------------------
        # 1. KUNCI PESANAN
        # -----------------------------------------------------

        cursor.execute("""
            SELECT id_pesanan,
                   no_pesanan

            FROM pesanan

            WHERE id_pesanan = %s

            FOR UPDATE
        """, (id_pesanan,))

        data_pesanan = cursor.fetchone()

        if not data_pesanan:
            raise Exception(
                "Pesanan tidak ditemukan."
            )

        no_pesanan = data_pesanan["no_pesanan"]

        # -----------------------------------------------------
        # 2. AMBIL DETAIL LAMA
        # -----------------------------------------------------

        cursor.execute("""
            SELECT id_varian,
                   jumlah_pesanan

            FROM detail_pesanan

            WHERE id_pesanan = %s

            FOR UPDATE
        """, (id_pesanan,))

        detail_lama = cursor.fetchall()

        # -----------------------------------------------------
        # 3. KEMBALIKAN STOK LAMA
        # -----------------------------------------------------

        for item in detail_lama:

            cursor.execute("""
                UPDATE stok_barang

                SET stok = stok + %s

                WHERE id_varian = %s
            """, (
                item["jumlah_pesanan"],
                item["id_varian"]
            ))

            if cursor.rowcount != 1:
                raise Exception(
                    f"Data stok varian "
                    f"{item['id_varian']} tidak ditemukan."
                )

        # -----------------------------------------------------
        # 4. CEK DATA BARU
        # -----------------------------------------------------

        varian_dipakai = set()
        data_barang_baru = []
        total_harga_baru = 0

        for id_varian_text, jumlah_text in zip(
            id_varian_list,
            jumlah_list
        ):

            id_varian = int(id_varian_text)
            jumlah = int(jumlah_text)

            if id_varian in varian_dipakai:
                raise Exception(
                    "Varian produk yang sama tidak boleh "
                    "dipilih dua kali."
                )

            varian_dipakai.add(id_varian)

            if jumlah <= 0:
                raise Exception(
                    "Jumlah pesanan harus lebih dari 0."
                )

            cursor.execute("""
                SELECT p.harga_jual,
                       sb.stok

                FROM varian_produk vp

                JOIN produk p
                    ON vp.id_produk = p.id_produk

                JOIN stok_barang sb
                    ON vp.id_varian = sb.id_varian

                WHERE vp.id_varian = %s

                FOR UPDATE
            """, (id_varian,))

            data_varian = cursor.fetchone()

            if not data_varian:
                raise Exception(
                    "Varian produk atau stok tidak ditemukan."
                )

            harga_jual = float(
                data_varian["harga_jual"] or 0
            )

            stok_tersedia = int(
                data_varian["stok"] or 0
            )

            if jumlah > stok_tersedia:
                raise Exception(
                    f"Stok tidak cukup. "
                    f"Stok tersedia: {stok_tersedia}, "
                    f"jumlah yang diminta: {jumlah}."
                )

            subtotal = harga_jual * jumlah

            total_harga_baru += subtotal

            data_barang_baru.append({
                "id_varian": id_varian,
                "jumlah": jumlah,
                "harga_jual": harga_jual,
                "subtotal": subtotal
            })

        # -----------------------------------------------------
        # 5. HAPUS BARANG KELUAR LAMA
        # -----------------------------------------------------

        cursor.execute("""
            DELETE FROM barang_keluar

            WHERE alasan = 'Penjualan'
              AND keterangan = %s
        """, (
            f"Pesanan {no_pesanan}",
        ))

        # -----------------------------------------------------
        # 6. HAPUS DETAIL LAMA
        # -----------------------------------------------------

        cursor.execute("""
            DELETE FROM detail_pesanan

            WHERE id_pesanan = %s
        """, (
            id_pesanan,
        ))

        # -----------------------------------------------------
        # 7. SIMPAN DETAIL BARU + BARANG KELUAR
        # -----------------------------------------------------

        for item in data_barang_baru:

            id_varian = item["id_varian"]
            jumlah = item["jumlah"]
            harga_jual = item["harga_jual"]
            subtotal = item["subtotal"]

            cursor.execute("""
                INSERT INTO detail_pesanan
                (
                    id_pesanan,
                    id_varian,
                    jumlah_pesanan,
                    harga_jual,
                    subtotal
                )

                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                id_pesanan,
                id_varian,
                jumlah,
                harga_jual,
                subtotal
            ))

            cursor.execute("""
                INSERT INTO barang_keluar
                (
                    id_varian,
                    jumlah,
                    alasan,
                    tanggal,
                    keterangan
                )

                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                id_varian,
                jumlah,
                "Penjualan",
                tanggal.split("T")[0],
                f"Pesanan {no_pesanan}"
            ))

            cursor.execute("""
                UPDATE stok_barang

                SET stok = stok - %s

                WHERE id_varian = %s
                  AND stok >= %s
            """, (
                jumlah,
                id_varian,
                jumlah
            ))

            if cursor.rowcount != 1:
                raise Exception(
                    "Stok tidak mencukupi saat "
                    "menyimpan perubahan."
                )

        # -----------------------------------------------------
        # 8. UPDATE PESANAN
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE pesanan

            SET no_hp = %s,
                nama_customor = %s,
                alamat = %s,
                tanggal = %s,
                STATUS = %s,
                total_harga = %s,
                metode_pengambilan = %s,
                keterangan = %s

            WHERE id_pesanan = %s
        """, (
            no_hp,
            nama_customer,
            alamat,
            tanggal,
            status,
            total_harga_baru,
            metode_pengambilan,
            keterangan,
            id_pesanan
        ))

        # -----------------------------------------------------
        # 9. COMMIT
        # -----------------------------------------------------

        conn.commit()

    except Exception as e:

        conn.rollback()

        print(
            "ERROR EDIT PESANAN:",
            e
        )

        return (
            f"Edit pesanan gagal: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/pesanan")


# =========================================================
# BARANG MASUK
# =========================================================

@app.route("/barang-masuk")
def barang_masuk():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        cursor.execute("""
            SELECT COALESCE(
                SUM(dbm.jumlah),
                0
            ) AS total_barang_masuk

            FROM detail_barang_masuk dbm
        """)

        total_barang_masuk = cursor.fetchone()[
            "total_barang_masuk"
        ]

        cursor.execute("""
            SELECT COALESCE(
                SUM(dbm.jumlah),
                0
            ) AS barang_masuk_hari_ini

            FROM detail_barang_masuk dbm

            JOIN barang_masuk bm
                ON dbm.id_barang_masuk =
                   bm.id_barang_masuk

            WHERE DATE(bm.tanggal) = CURDATE()
        """)

        barang_masuk_hari_ini = cursor.fetchone()[
            "barang_masuk_hari_ini"
        ]

        cursor.execute("""
            SELECT COUNT(*) AS total_transaksi

            FROM barang_masuk
        """)

        total_transaksi = cursor.fetchone()[
            "total_transaksi"
        ]

        cursor.execute("""
            SELECT COALESCE(
                SUM(
                    jumlah * harga_beli
                ),
                0
            ) AS total_nilai_barang

            FROM detail_barang_masuk
        """)

        total_nilai_barang = cursor.fetchone()[
            "total_nilai_barang"
        ]

        cursor.execute("""
            SELECT bm.id_barang_masuk,
                   bm.no_faktur,
                   bm.tanggal,
                   bm.nama_penjahit,
                   p.nama_produk,
                   w.nama_warna,
                   dbm.id_varian,
                   SUM(dbm.jumlah) AS jumlah,
                   SUM(
                       dbm.jumlah *
                       dbm.harga_beli
                   ) AS total

            FROM barang_masuk bm

            JOIN detail_barang_masuk dbm
                ON bm.id_barang_masuk =
                   dbm.id_barang_masuk

            JOIN varian_produk vp
                ON dbm.id_varian =
                   vp.id_varian

            JOIN produk p
                ON vp.id_produk =
                   p.id_produk

            JOIN warna w
                ON vp.id_warna =
                   w.id_warna

            GROUP BY
                bm.id_barang_masuk,
                bm.no_faktur,
                bm.tanggal,
                bm.nama_penjahit,
                p.nama_produk,
                w.nama_warna,
                dbm.id_varian

            ORDER BY bm.tanggal DESC
        """)

        data_barang_masuk = cursor.fetchall()

        cursor.execute("""
            SELECT vp.id_varian,
                   p.nama_produk,
                   w.nama_warna

            FROM varian_produk vp

            JOIN produk p
                ON vp.id_produk =
                   p.id_produk

            JOIN warna w
                ON vp.id_warna =
                   w.id_warna

            ORDER BY p.nama_produk,
                     w.nama_warna
        """)

        data_varian = cursor.fetchall()

        return render_template(
            "barang_masuk.html",
            total_barang_masuk=total_barang_masuk,
            barang_masuk_hari_ini=barang_masuk_hari_ini,
            total_transaksi=total_transaksi,
            total_nilai_barang=total_nilai_barang,
            barang_masuk=data_barang_masuk,
            varian=data_varian
        )

    finally:
        cursor.close()
        conn.close()


# =========================================================
# TAMBAH BARANG MASUK
# =========================================================

@app.route("/barang-masuk/tambah", methods=["POST"])
def tambah_barang_masuk():

    nama_penjahit = request.form["nama_penjahit"]
    tanggal = request.form["tanggal"]
    id_varian = int(request.form["id_varian"])
    jumlah = int(request.form["jumlah"])
    harga_beli = request.form["harga_beli"]
    keterangan = request.form.get(
        "keterangan",
        ""
    )

    conn = get_db_connection()
    cursor = conn.cursor()

    try:

        # -----------------------------------------------------
        # VALIDASI JUMLAH
        # -----------------------------------------------------

        if jumlah <= 0:
            raise Exception(
                "Jumlah barang masuk harus lebih dari 0."
            )

        # -----------------------------------------------------
        # KUNCI STOK
        # -----------------------------------------------------

        cursor.execute("""
            SELECT stok

            FROM stok_barang

            WHERE id_varian = %s

            FOR UPDATE
        """, (
            id_varian,
        ))

        data_stok = cursor.fetchone()

        if not data_stok:
            raise Exception(
                "Data stok varian tidak ditemukan."
            )

        # -----------------------------------------------------
        # GENERATE NO FAKTUR
        # -----------------------------------------------------

        cursor.execute("""
            SELECT no_faktur

            FROM barang_masuk

            WHERE no_faktur LIKE 'BM-%'

            ORDER BY id_barang_masuk DESC

            LIMIT 1
        """)

        data_terakhir = cursor.fetchone()

        if data_terakhir:

            nomor_baru = (
                int(
                    data_terakhir[0]
                    .replace("BM-", "")
                )
                + 1
            )

        else:
            nomor_baru = 1

        no_faktur = f"BM-{nomor_baru:04d}"

        # -----------------------------------------------------
        # INSERT HEADER
        # -----------------------------------------------------

        cursor.execute("""
            INSERT INTO barang_masuk
            (
                no_faktur,
                nama_penjahit,
                tanggal,
                keterangan
            )

            VALUES
            (
                %s,
                %s,
                %s,
                %s
            )
        """, (
            no_faktur,
            nama_penjahit,
            tanggal,
            keterangan
        ))

        id_barang_masuk = cursor.lastrowid

        # -----------------------------------------------------
        # INSERT DETAIL
        # -----------------------------------------------------

        cursor.execute("""
            INSERT INTO detail_barang_masuk
            (
                id_barang_masuk,
                id_varian,
                jumlah,
                harga_beli
            )

            VALUES
            (
                %s,
                %s,
                %s,
                %s
            )
        """, (
            id_barang_masuk,
            id_varian,
            jumlah,
            harga_beli
        ))

        # -----------------------------------------------------
        # TAMBAH STOK
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE stok_barang

            SET stok = stok + %s

            WHERE id_varian = %s
        """, (
            jumlah,
            id_varian
        ))

        if cursor.rowcount != 1:
            raise Exception(
                "Stok varian tidak ditemukan."
            )

        conn.commit()

    except Exception as e:

        conn.rollback()

        print(
            "ERROR TAMBAH BARANG MASUK:",
            e
        )

        return (
            f"Tambah barang masuk gagal: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/barang-masuk")


# =========================================================
# HAPUS BARANG MASUK
# =========================================================

@app.route(
    "/barang-masuk/hapus/<int:id_barang_masuk>",
    methods=["POST"]
)
def hapus_barang_masuk(id_barang_masuk):

    conn = get_db_connection()
    cursor = conn.cursor()

    try:

        # -----------------------------------------------------
        # AMBIL DETAIL
        # -----------------------------------------------------

        cursor.execute("""
            SELECT id_varian,
                   jumlah

            FROM detail_barang_masuk

            WHERE id_barang_masuk = %s

            FOR UPDATE
        """, (
            id_barang_masuk,
        ))

        data = cursor.fetchone()

        if not data:
            raise Exception(
                "Data barang masuk tidak ditemukan."
            )

        id_varian = data[0]
        jumlah = int(data[1])

        # -----------------------------------------------------
        # KUNCI STOK
        # -----------------------------------------------------

        cursor.execute("""
            SELECT stok

            FROM stok_barang

            WHERE id_varian = %s

            FOR UPDATE
        """, (
            id_varian,
        ))

        data_stok = cursor.fetchone()

        if not data_stok:
            raise Exception(
                "Data stok tidak ditemukan."
            )

        stok_sekarang = int(
            data_stok[0]
        )

        # -----------------------------------------------------
        # CEK STOK
        # -----------------------------------------------------

        if jumlah > stok_sekarang:
            raise Exception(
                f"Transaksi tidak bisa dihapus "
                f"karena stok saat ini tidak mencukupi. "
                f"Stok tersedia: {stok_sekarang}, "
                f"jumlah yang akan dikurangi: {jumlah}."
            )

        # -----------------------------------------------------
        # KURANGI STOK
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE stok_barang

            SET stok = stok - %s

            WHERE id_varian = %s
              AND stok >= %s
        """, (
            jumlah,
            id_varian,
            jumlah
        ))

        if cursor.rowcount != 1:
            raise Exception(
                "Stok tidak mencukupi."
            )

        # -----------------------------------------------------
        # HAPUS DETAIL
        # -----------------------------------------------------

        cursor.execute("""
            DELETE FROM detail_barang_masuk

            WHERE id_barang_masuk = %s
        """, (
            id_barang_masuk,
        ))

        # -----------------------------------------------------
        # HAPUS HEADER
        # -----------------------------------------------------

        cursor.execute("""
            DELETE FROM barang_masuk

            WHERE id_barang_masuk = %s
        """, (
            id_barang_masuk,
        ))

        conn.commit()

    except Exception as e:

        conn.rollback()

        print(
            "ERROR HAPUS BARANG MASUK:",
            e
        )

        return (
            f"Hapus barang masuk gagal: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/barang-masuk")


# =========================================================
# EDIT BARANG MASUK
# =========================================================

@app.route(
    "/barang-masuk/edit/<int:id_barang_masuk>",
    methods=["POST"]
)
def edit_barang_masuk(id_barang_masuk):

    nama_penjahit = request.form[
        "nama_penjahit"
    ]

    tanggal = request.form[
        "tanggal"
    ]

    id_varian_baru = int(
        request.form["id_varian"]
    )

    jumlah_baru = int(
        request.form["jumlah"]
    )

    harga_beli = request.form[
        "harga_beli"
    ]

    keterangan = request.form.get(
        "keterangan",
        ""
    )

    conn = get_db_connection()
    cursor = conn.cursor()

    try:

        # -----------------------------------------------------
        # VALIDASI
        # -----------------------------------------------------

        if jumlah_baru <= 0:
            raise Exception(
                "Jumlah harus lebih dari 0."
            )

        # -----------------------------------------------------
        # AMBIL DATA LAMA
        # -----------------------------------------------------

        cursor.execute("""
            SELECT id_varian,
                   jumlah

            FROM detail_barang_masuk

            WHERE id_barang_masuk = %s

            FOR UPDATE
        """, (
            id_barang_masuk,
        ))

        data_lama = cursor.fetchone()

        if not data_lama:
            raise Exception(
                "Data barang masuk tidak ditemukan."
            )

        id_varian_lama = int(
            data_lama[0]
        )

        jumlah_lama = int(
            data_lama[1]
        )

        # =====================================================
        # VARIAN TIDAK BERUBAH
        # =====================================================

        if id_varian_lama == id_varian_baru:

            selisih = (
                jumlah_baru -
                jumlah_lama
            )

            # -------------------------------------------------
            # JUMLAH BERTAMBAH
            # -------------------------------------------------

            if selisih > 0:

                cursor.execute("""
                    SELECT stok

                    FROM stok_barang

                    WHERE id_varian = %s

                    FOR UPDATE
                """, (
                    id_varian_baru,
                ))

                data_stok = cursor.fetchone()

                if not data_stok:
                    raise Exception(
                        "Data stok tidak ditemukan."
                    )

                cursor.execute("""
                    UPDATE stok_barang

                    SET stok = stok + %s

                    WHERE id_varian = %s
                """, (
                    selisih,
                    id_varian_baru
                ))

            # -------------------------------------------------
            # JUMLAH BERKURANG
            # -------------------------------------------------

            elif selisih < 0:

                jumlah_dikurangi = abs(
                    selisih
                )

                cursor.execute("""
                    SELECT stok

                    FROM stok_barang

                    WHERE id_varian = %s

                    FOR UPDATE
                """, (
                    id_varian_lama,
                ))

                data_stok = cursor.fetchone()

                if not data_stok:
                    raise Exception(
                        "Data stok tidak ditemukan."
                    )

                stok_sekarang = int(
                    data_stok[0]
                )

                if jumlah_dikurangi > stok_sekarang:
                    raise Exception(
                        f"Stok tidak cukup untuk "
                        f"mengurangi "
                        f"{jumlah_dikurangi} barang. "
                        f"Stok tersedia: "
                        f"{stok_sekarang}"
                    )

                cursor.execute("""
                    UPDATE stok_barang

                    SET stok = stok - %s

                    WHERE id_varian = %s
                      AND stok >= %s
                """, (
                    jumlah_dikurangi,
                    id_varian_lama,
                    jumlah_dikurangi
                ))

                if cursor.rowcount != 1:
                    raise Exception(
                        "Stok tidak mencukupi."
                    )

        # =====================================================
        # VARIAN BERUBAH
        # =====================================================

        else:

            # -------------------------------------------------
            # KEMBALIKAN STOK VARIAN LAMA
            # -------------------------------------------------

            cursor.execute("""
                UPDATE stok_barang

                SET stok = stok + %s

                WHERE id_varian = %s
            """, (
                jumlah_lama,
                id_varian_lama
            ))

            if cursor.rowcount != 1:
                raise Exception(
                    "Stok varian lama "
                    "tidak ditemukan."
                )

            # -------------------------------------------------
            # KUNCI STOK VARIAN BARU
            # -------------------------------------------------

            cursor.execute("""
                SELECT stok

                FROM stok_barang

                WHERE id_varian = %s

                FOR UPDATE
            """, (
                id_varian_baru,
            ))

            data_stok_baru = cursor.fetchone()

            if not data_stok_baru:
                raise Exception(
                    "Stok varian baru "
                    "tidak ditemukan."
                )

            stok_baru = int(
                data_stok_baru[0]
            )

            # -------------------------------------------------
            # CEK STOK VARIAN BARU
            # -------------------------------------------------

            if jumlah_baru > stok_baru:
                raise Exception(
                    f"Stok varian baru tidak cukup. "
                    f"Stok tersedia: {stok_baru}"
                )

            # -------------------------------------------------
            # KURANGI STOK VARIAN BARU
            # -------------------------------------------------

            cursor.execute("""
                UPDATE stok_barang

                SET stok = stok - %s

                WHERE id_varian = %s
                  AND stok >= %s
            """, (
                jumlah_baru,
                id_varian_baru,
                jumlah_baru
            ))

            if cursor.rowcount != 1:
                raise Exception(
                    "Stok varian baru "
                    "tidak mencukupi."
                )

        # -----------------------------------------------------
        # UPDATE DETAIL
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE detail_barang_masuk

            SET id_varian = %s,
                jumlah = %s,
                harga_beli = %s

            WHERE id_barang_masuk = %s
        """, (
            id_varian_baru,
            jumlah_baru,
            harga_beli,
            id_barang_masuk
        ))

        # -----------------------------------------------------
        # UPDATE HEADER
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE barang_masuk

            SET nama_penjahit = %s,
                tanggal = %s,
                keterangan = %s

            WHERE id_barang_masuk = %s
        """, (
            nama_penjahit,
            tanggal,
            keterangan,
            id_barang_masuk
        ))

        conn.commit()

    except Exception as e:

        conn.rollback()

        print(
            "ERROR EDIT BARANG MASUK:",
            e
        )

        return (
            f"Edit barang masuk gagal: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/barang-masuk")


# =========================================================
# BARANG KELUAR
# =========================================================

@app.route("/barang-keluar")
def barang_keluar():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        cursor.execute("""
            SELECT COALESCE(
                SUM(jumlah),
                0
            ) AS total_barang_keluar

            FROM barang_keluar
        """)

        total_barang_keluar = cursor.fetchone()[
            "total_barang_keluar"
        ]

        cursor.execute("""
            SELECT COALESCE(
                SUM(jumlah),
                0
            ) AS barang_keluar_hari_ini

            FROM barang_keluar

            WHERE DATE(tanggal) = CURDATE()
        """)

        barang_keluar_hari_ini = cursor.fetchone()[
            "barang_keluar_hari_ini"
        ]

        cursor.execute("""
            SELECT COUNT(*) AS total_transaksi_keluar

            FROM barang_keluar
        """)

        total_transaksi_keluar = cursor.fetchone()[
            "total_transaksi_keluar"
        ]

        cursor.execute("""
            SELECT COALESCE(
                SUM(jumlah),
                0
            ) AS barang_rusak

            FROM barang_keluar

            WHERE alasan = 'Rusak'
        """)

        barang_rusak = cursor.fetchone()[
            "barang_rusak"
        ]

        cursor.execute("""
            SELECT bk.id_barang_keluar,
                   bk.id_varian,
                   bk.jumlah,
                   bk.alasan,
                   bk.tanggal,
                   bk.keterangan,
                   p.nama_produk,
                   w.nama_warna

            FROM barang_keluar bk

            JOIN varian_produk vp
                ON bk.id_varian =
                   vp.id_varian

            JOIN produk p
                ON vp.id_produk =
                   p.id_produk

            JOIN warna w
                ON vp.id_warna =
                   w.id_warna

            ORDER BY
                bk.tanggal DESC,
                bk.id_barang_keluar DESC
        """)

        data_barang_keluar = cursor.fetchall()

        cursor.execute("""
            SELECT vp.id_varian,
                   p.nama_produk,
                   w.nama_warna

            FROM varian_produk vp

            JOIN produk p
                ON vp.id_produk =
                   p.id_produk

            JOIN warna w
                ON vp.id_warna =
                   w.id_warna

            ORDER BY p.nama_produk,
                     w.nama_warna
        """)

        data_varian = cursor.fetchall()

        return render_template(
            "barang_keluar.html",
            total_barang_keluar=total_barang_keluar,
            barang_keluar_hari_ini=barang_keluar_hari_ini,
            total_transaksi_keluar=total_transaksi_keluar,
            barang_rusak=barang_rusak,
            barang_keluar=data_barang_keluar,
            varian=data_varian
        )

    finally:
        cursor.close()
        conn.close()


# =========================================================
# TAMBAH BARANG KELUAR
# =========================================================

@app.route(
    "/barang-keluar/tambah",
    methods=["POST"]
)
def tambah_barang_keluar():

    tanggal = request.form["tanggal"]

    id_varian = int(
        request.form["id_varian"]
    )

    jumlah = int(
        request.form["jumlah"]
    )

    alasan = request.form["alasan"]

    keterangan = request.form.get(
        "keterangan",
        ""
    )

    conn = get_db_connection()
    cursor = conn.cursor()

    try:

        if jumlah <= 0:
            raise Exception(
                "Jumlah harus lebih dari 0."
            )

        # -----------------------------------------------------
        # KUNCI STOK
        # -----------------------------------------------------

        cursor.execute("""
            SELECT stok

            FROM stok_barang

            WHERE id_varian = %s

            FOR UPDATE
        """, (
            id_varian,
        ))

        data_stok = cursor.fetchone()

        if not data_stok:
            raise Exception(
                "Data stok tidak ditemukan."
            )

        stok_tersedia = int(
            data_stok[0]
        )

        if jumlah > stok_tersedia:
            raise Exception(
                f"Stok tidak cukup. "
                f"Stok tersedia: {stok_tersedia}"
            )

        # -----------------------------------------------------
        # INSERT BARANG KELUAR
        # -----------------------------------------------------

        cursor.execute("""
            INSERT INTO barang_keluar
            (
                id_varian,
                jumlah,
                alasan,
                tanggal,
                keterangan
            )

            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s
            )
        """, (
            id_varian,
            jumlah,
            alasan,
            tanggal.split("T")[0],
            keterangan
        ))

        # -----------------------------------------------------
        # KURANGI STOK
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE stok_barang

            SET stok = stok - %s

            WHERE id_varian = %s
              AND stok >= %s
        """, (
            jumlah,
            id_varian,
            jumlah
        ))

        if cursor.rowcount != 1:
            raise Exception(
                "Stok tidak mencukupi."
            )

        conn.commit()

    except Exception as e:

        conn.rollback()

        print(
            "ERROR TAMBAH BARANG KELUAR:",
            e
        )

        return (
            f"Tambah barang keluar gagal: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/barang-keluar")


# =========================================================
# HAPUS BARANG KELUAR
# =========================================================

@app.route(
    "/barang-keluar/hapus/<int:id_barang_keluar>",
    methods=["POST"]
)
def hapus_barang_keluar(id_barang_keluar):

    conn = get_db_connection()
    cursor = conn.cursor()

    try:

        cursor.execute("""
            SELECT id_varian,
                   jumlah

            FROM barang_keluar

            WHERE id_barang_keluar = %s

            FOR UPDATE
        """, (
            id_barang_keluar,
        ))

        data = cursor.fetchone()

        if not data:
            raise Exception(
                "Data barang keluar tidak ditemukan."
            )

        id_varian = data[0]
        jumlah = int(data[1])

        # -----------------------------------------------------
        # KEMBALIKAN STOK
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE stok_barang

            SET stok = stok + %s

            WHERE id_varian = %s
        """, (
            jumlah,
            id_varian
        ))

        if cursor.rowcount != 1:
            raise Exception(
                "Data stok tidak ditemukan."
            )

        # -----------------------------------------------------
        # HAPUS DATA
        # -----------------------------------------------------

        cursor.execute("""
            DELETE FROM barang_keluar

            WHERE id_barang_keluar = %s
        """, (
            id_barang_keluar,
        ))

        conn.commit()

    except Exception as e:

        conn.rollback()

        print(
            "ERROR HAPUS BARANG KELUAR:",
            e
        )

        return (
            f"Hapus barang keluar gagal: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/barang-keluar")


# =========================================================
# EDIT BARANG KELUAR
# =========================================================

@app.route(
    "/barang-keluar/edit/<int:id_barang_keluar>",
    methods=["POST"]
)
def edit_barang_keluar(id_barang_keluar):

    tanggal = request.form["tanggal"]

    id_varian_baru = int(
        request.form["id_varian"]
    )

    jumlah_baru = int(
        request.form["jumlah"]
    )

    alasan = request.form["alasan"]

    keterangan = request.form.get(
        "keterangan",
        ""
    )

    conn = get_db_connection()
    cursor = conn.cursor()

    try:

        if jumlah_baru <= 0:
            raise Exception(
                "Jumlah harus lebih dari 0."
            )

        # -----------------------------------------------------
        # AMBIL DATA LAMA
        # -----------------------------------------------------

        cursor.execute("""
            SELECT id_varian,
                   jumlah

            FROM barang_keluar

            WHERE id_barang_keluar = %s

            FOR UPDATE
        """, (
            id_barang_keluar,
        ))

        data_lama = cursor.fetchone()

        if not data_lama:
            raise Exception(
                "Data barang keluar tidak ditemukan."
            )

        id_varian_lama = int(
            data_lama[0]
        )

        jumlah_lama = int(
            data_lama[1]
        )

        # -----------------------------------------------------
        # KEMBALIKAN STOK LAMA
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE stok_barang

            SET stok = stok + %s

            WHERE id_varian = %s
        """, (
            jumlah_lama,
            id_varian_lama
        ))

        if cursor.rowcount != 1:
            raise Exception(
                "Stok varian lama tidak ditemukan."
            )

        # -----------------------------------------------------
        # KUNCI STOK BARU
        # -----------------------------------------------------

        cursor.execute("""
            SELECT stok

            FROM stok_barang

            WHERE id_varian = %s

            FOR UPDATE
        """, (
            id_varian_baru,
        ))

        data_stok = cursor.fetchone()

        if not data_stok:
            raise Exception(
                "Data stok tidak ditemukan."
            )

        stok_tersedia = int(
            data_stok[0]
        )

        if jumlah_baru > stok_tersedia:
            raise Exception(
                f"Stok tidak cukup. "
                f"Stok tersedia: {stok_tersedia}"
            )

        # -----------------------------------------------------
        # UPDATE DATA
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE barang_keluar

            SET id_varian = %s,
                jumlah = %s,
                alasan = %s,
                tanggal = %s,
                keterangan = %s

            WHERE id_barang_keluar = %s
        """, (
            id_varian_baru,
            jumlah_baru,
            alasan,
            tanggal.split("T")[0],
            keterangan,
            id_barang_keluar
        ))

        # -----------------------------------------------------
        # KURANGI STOK BARU
        # -----------------------------------------------------

        cursor.execute("""
            UPDATE stok_barang

            SET stok = stok - %s

            WHERE id_varian = %s
              AND stok >= %s
        """, (
            jumlah_baru,
            id_varian_baru,
            jumlah_baru
        ))

        if cursor.rowcount != 1:
            raise Exception(
                "Stok tidak mencukupi."
            )

        conn.commit()

    except Exception as e:

        conn.rollback()

        print(
            "ERROR EDIT BARANG KELUAR:",
            e
        )

        return (
            f"Edit barang keluar gagal: {e}",
            400
        )

    finally:
        cursor.close()
        conn.close()

    return redirect("/barang-keluar")


# =========================================================
# LAPORAN
# =========================================================

@app.route("/laporan")
def laporan():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        cursor.execute("""
            SELECT COUNT(*) AS total_pesanan,

                   COALESCE(
                       SUM(
                           CASE
                               WHEN STATUS != 'dibatalkan'
                               THEN total_harga
                               ELSE 0
                           END
                       ),
                       0
                   ) AS total_penjualan

            FROM pesanan
        """)

        ringkasan_pesanan = cursor.fetchone()

        cursor.execute("""
            SELECT COALESCE(
                SUM(jumlah),
                0
            ) AS total_barang_masuk

            FROM detail_barang_masuk
        """)

        ringkasan_masuk = cursor.fetchone()

        cursor.execute("""
            SELECT COALESCE(
                SUM(jumlah),
                0
            ) AS total_barang_keluar

            FROM barang_keluar
        """)

        ringkasan_keluar = cursor.fetchone()

        cursor.execute("""
            SELECT id_pesanan,
                   no_pesanan,
                   tanggal,
                   nama_customor,
                   no_hp,
                   total_harga,
                   STATUS,
                   metode_pengambilan

            FROM pesanan

            ORDER BY tanggal DESC
        """)

        data_laporan = cursor.fetchall()

        cursor.execute("""
            SELECT DATE_FORMAT(
                       tanggal,
                       '%Y-%m-%d'
                   ) AS tanggal,

                   COALESCE(
                       SUM(total_harga),
                       0
                   ) AS penjualan

            FROM pesanan

            WHERE STATUS != 'dibatalkan'

            GROUP BY DATE(tanggal)

            ORDER BY tanggal ASC
        """)

        grafik_penjualan = cursor.fetchall()

        cursor.execute("""
            SELECT DATE_FORMAT(
                       bm.tanggal,
                       '%Y-%m-%d'
                   ) AS tanggal,

                   COALESCE(
                       SUM(dbm.jumlah),
                       0
                   ) AS jumlah

            FROM barang_masuk bm

            JOIN detail_barang_masuk dbm
                ON bm.id_barang_masuk =
                   dbm.id_barang_masuk

            GROUP BY DATE(bm.tanggal)

            ORDER BY tanggal ASC
        """)

        grafik_masuk = cursor.fetchall()

        cursor.execute("""
            SELECT DATE_FORMAT(
                       tanggal,
                       '%Y-%m-%d'
                   ) AS tanggal,

                   COALESCE(
                       SUM(jumlah),
                       0
                   ) AS jumlah

            FROM barang_keluar

            GROUP BY tanggal

            ORDER BY tanggal ASC
        """)

        grafik_keluar = cursor.fetchall()

        return render_template(
            "laporan.html",
            ringkasan_pesanan=ringkasan_pesanan,
            ringkasan_masuk=ringkasan_masuk,
            ringkasan_keluar=ringkasan_keluar,
            data_laporan=data_laporan,
            grafik_penjualan=grafik_penjualan,
            grafik_masuk=grafik_masuk,
            grafik_keluar=grafik_keluar
        )

    finally:
        cursor.close()
        conn.close()


# =========================================================
# EXPORT LAPORAN EXCEL
# =========================================================

@app.route("/laporan/export")
def export_laporan():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        cursor.execute("""
            SELECT id_pesanan,
                   no_pesanan,
                   tanggal,
                   nama_customor,
                   no_hp,
                   total_harga,
                   STATUS,
                   metode_pengambilan

            FROM pesanan

            ORDER BY tanggal DESC
        """)

        data_laporan = cursor.fetchall()

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Laporan Penjualan"

        sheet["A1"] = (
            "LAPORAN PENJUALAN YN HIJAB"
        )

        sheet["A1"].font = Font(
            bold=True,
            size=16
        )

        sheet["A1"].alignment = Alignment(
            horizontal="center"
        )

        sheet.merge_cells("A1:G1")

        header = [
            "No Pesanan",
            "Tanggal",
            "Pelanggan",
            "No HP",
            "Total",
            "Status",
            "Metode Pengambilan"
        ]

        for column, value in enumerate(
            header,
            start=1
        ):

            cell = sheet.cell(
                row=3,
                column=column
            )

            cell.value = value
            cell.font = Font(
                bold=True
            )

        row_number = 4

        for item in data_laporan:

            sheet.cell(
                row=row_number,
                column=1
            ).value = item[
                "no_pesanan"
            ]

            sheet.cell(
                row=row_number,
                column=2
            ).value = item[
                "tanggal"
            ]

            sheet.cell(
                row=row_number,
                column=3
            ).value = item[
                "nama_customor"
            ]

            sheet.cell(
                row=row_number,
                column=4
            ).value = item[
                "no_hp"
            ]

            sheet.cell(
                row=row_number,
                column=5
            ).value = float(
                item["total_harga"] or 0
            )

            sheet.cell(
                row=row_number,
                column=6
            ).value = item[
                "STATUS"
            ]

            sheet.cell(
                row=row_number,
                column=7
            ).value = item[
                "metode_pengambilan"
            ]

            row_number += 1

        for row in range(
            4,
            row_number
        ):

            sheet.cell(
                row=row,
                column=5
            ).number_format = "#,##0"

        widths = {
            "A": 15,
            "B": 22,
            "C": 25,
            "D": 20,
            "E": 18,
            "F": 20,
            "G": 25
        }

        for column, width in widths.items():
            sheet.column_dimensions[
                column
            ].width = width

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=(
                "laporan_yn_hijab.xlsx"
            ),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    finally:
        cursor.close()
        conn.close()


# =========================================================
# STOK
# =========================================================

@app.route("/stok")
def stok():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        cursor.execute("""
            SELECT sb.id_stok,
                   p.nama_produk,
                   p.kode_produk,
                   w.nama_warna,
                   r.kode_rak,
                   r.nama_rak,
                   sb.stok,
                   sb.stok_minimum

            FROM stok_barang sb

            JOIN varian_produk vp
                ON sb.id_varian =
                   vp.id_varian

            JOIN produk p
                ON vp.id_produk =
                   p.id_produk

            JOIN warna w
                ON vp.id_warna =
                   w.id_warna

            JOIN rak r
                ON sb.id_rak =
                   r.id_rak

            ORDER BY p.nama_produk ASC
        """)

        data_stok = cursor.fetchall()

        cursor.execute("""
            SELECT COUNT(*) AS total_varian

            FROM stok_barang
        """)

        total_varian = cursor.fetchone()[
            "total_varian"
        ]

        cursor.execute("""
            SELECT COALESCE(
                SUM(stok),
                0
            ) AS total_stok

            FROM stok_barang
        """)

        total_stok = cursor.fetchone()[
            "total_stok"
        ]

        cursor.execute("""
            SELECT COUNT(*) AS stok_menipis

            FROM stok_barang

            WHERE stok <= stok_minimum
        """)

        stok_menipis = cursor.fetchone()[
            "stok_menipis"
        ]

        return render_template(
            "stok.html",
            stok=data_stok,
            total_varian=total_varian,
            total_stok=total_stok,
            stok_menipis=stok_menipis
        )

    finally:
        cursor.close()
        conn.close()


# =========================================================
# JALANKAN FLASK
# =========================================================

if __name__ == "__main__":
    app.run(debug=True)